"""Parser for the Business Expert Press (BEP) transmittal workbook."""
from __future__ import annotations

import openpyxl

from app.domains.projects.po_intake import normalize as n

def _find_cell_by_label(ws, label: str) -> object:
    label_lower = label.casefold()
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().casefold() == label_lower:
                return ws.cell(row=cell.row, column=cell.column + 1).value
    return None

def parse(xlsx_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    warnings: list[str] = []
    
    ws = wb["Specs"] if "Specs" in wb.sheetnames else wb.active

    title = n.clean(_find_cell_by_label(ws, "Title"))
    subtitle = n.clean(_find_cell_by_label(ws, "Subtitle"))
    if title and subtitle:
        project_title = f"{title}: {subtitle}"
    else:
        project_title = title or subtitle

    author_raw = n.clean(_find_cell_by_label(ws, "Author name(s)"))
    author_names = [author_raw] if author_raw else []

    due_date_raw = _find_cell_by_label(ws, "Target final files posted")
    due_date, due_warning = n.parse_date_loose(due_date_raw)
    if due_warning:
        warnings.append(f"Target final files posted: {due_warning}")

    isbn_no = n.clean(_find_cell_by_label(ws, "Paperback ISBN"))
    if not isbn_no or isbn_no.upper() in ("TO COME", "TBA", "N/A"):
        isbn_no = n.clean(_find_cell_by_label(ws, "Hardcover ISBN"))
    if isbn_no and isbn_no.upper() in ("TO COME", "TBA", "N/A"):
        isbn_no = None

    isbn_val, isbn_warning = n.normalize_isbn(isbn_no)
    if isbn_warning:
        warnings.append(isbn_warning)

    ebook_isbn = n.clean(_find_cell_by_label(ws, "eISBN"))
    if ebook_isbn and ebook_isbn.upper() in ("TO COME", "TBA", "N/A"):
        ebook_isbn = None

    pages_raw = _find_cell_by_label(ws, "Total manuscript pages")
    manuscript_pages = n.to_int(pages_raw)

    chapters_raw = _find_cell_by_label(ws, "Number of chapters")
    chapter_count = n.to_int(chapters_raw)

    collection_editor = n.clean(_find_cell_by_label(ws, "Collection Editor"))

    ce_level_raw = n.clean(_find_cell_by_label(ws, "Level of copyedit"))
    copyediting_level = "Level 1"
    if ce_level_raw:
        ce_lower = ce_level_raw.lower()
        if "level 2" in ce_lower or "medium" in ce_lower:
            copyediting_level = "Level 2"
        elif "level 3" in ce_lower or "high" in ce_lower:
            copyediting_level = "Level 3"

    fields = {
        "project_title": project_title,
        "isbn_no": isbn_val,
        "due_date": due_date,
        "manuscript_pages": manuscript_pages,
        "chapter_count": chapter_count,
        "client_project_manager": collection_editor,
        "copyediting_level": copyediting_level,
        "author_names": author_names,
    }

    extras = {
        "ebook_isbn": ebook_isbn,
        "publisher": "Business Expert Press",
    }

    return {"fields": fields, "extras": extras, "warnings": warnings}
