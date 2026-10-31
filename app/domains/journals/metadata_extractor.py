import os
import re
import zipfile
import tempfile
from typing import Dict, Any, List, Optional
import docx


ARTICLE_TYPE_PATTERNS = [
    r'^(?:review|review\s+article|systematic\s+review(?:\s+and\s+meta-analysis)?|narrative\s+review|scoping\s+review|literature\s+review|mini[- ]review)$',
    r'^(?:original\s+article|original\s+research|research\s+article|original\s+paper|full\s+length\s+article|regular\s+article)$',
    r'^(?:case\s+report|case\s+study|clinical\s+case|case\s+series)$',
    r'^(?:short\s+communication|preliminary\s+communication|brief\s+communication|brief\s+report|rapid\s+communication)$',
    r'^(?:editorial|guest\s+editorial|editorial\s+note|editor\'s\s+note)$',
    r'^(?:commentary|expert\s+commentary|invited\s+commentary)$',
    r'^(?:letter\s+to\s+(?:the\s+)?editor|letter|correspondence|reply)$',
    r'^(?:special\s+report|special\s+article|special\s+issue\s+article)$',
    r'^(?:methodology|methods?\s+paper|study\s+protocol|protocol)$',
    r'^(?:clinical\s+trial|randomized\s+controlled\s+trial)$',
    r'^(?:perspective|viewpoint|opinion|essay|hypothesis)$',
]
ARTICLE_TYPE_REGEX = re.compile('|'.join(ARTICLE_TYPE_PATTERNS), re.IGNORECASE)


def is_article_type_heading(text: str) -> bool:
    """Returns True if the text represents an article type classification rather than a title."""
    t = text.strip()
    if not t:
        return False
    clean = re.sub(r'[:\.\-\–\—]+$', '', t).strip()
    if ARTICLE_TYPE_REGEX.match(clean):
        return True
    return clean.lower() in [
        'review', 'review article', 'original article', 'case report', 'case study',
        'short communication', 'editorial', 'commentary', 'letter to editor', 'letter to the editor',
        'special report', 'methodology', 'clinical trial', 'perspective', 'viewpoint', 'brief report'
    ]


def extract_docx_metadata_advanced(docx_path: str) -> Dict[str, Any]:
    """
    Automated DOCX Metadata Extractor Module.
    Parses manuscript DOCX to extract:
    - Article Title (validated against article type classification)
    - Article Type
    - DOI
    - Authors & Affiliations
    - Corresponding Email
    - Abstract
    - Keywords
    - Word count & Page estimations
    """
    metadata: Dict[str, Any] = {
        "article_title": None,
        "detected_article_type": None,
        "article_doi": None,
        "lead_author": None,
        "co_authors": [],
        "affiliations": [],
        "corresponding_email": None,
        "abstract": None,
        "keywords": [],
        "word_count": 0,
        "estimated_pages": 1
    }

    if not os.path.exists(docx_path):
        return metadata

    try:
        doc = docx.Document(docx_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if not paragraphs:
            return metadata

        # 1. Total Word Count & Page Estimation
        total_words = sum(len(p.split()) for p in paragraphs)
        metadata["word_count"] = total_words
        metadata["estimated_pages"] = max(1, round(total_words / 450))

        # 2. Extract DOI using Regex
        full_text = "\n".join(paragraphs)
        doi_match = re.search(r'10\.\d{4,9}/[-._;()/:A-Za-z0-9]+', full_text)
        if doi_match:
            metadata["article_doi"] = doi_match.group(0).rstrip('.,;')
        else:
            # Check text containing 'doi'
            doi_text_match = re.search(r'doi:\s*(10\.\S+)', full_text, re.IGNORECASE)
            if doi_text_match:
                metadata["article_doi"] = doi_text_match.group(1).rstrip('.,;')

        # 3. Extract Email using Regex
        email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', full_text)
        if email_match:
            metadata["corresponding_email"] = email_match.group(0)

        # 4. Extract Article Type & Real Article Title
        detected_article_type = None
        title_candidate = None
        title_para_idx = -1

        for idx, p in enumerate(paragraphs[:15]):
            if is_article_type_heading(p):
                if not detected_article_type:
                    detected_article_type = p
                continue
            if re.match(r'^(?:abstract|keywords?|introduction|doi:|https?://|www\.)\b', p, re.I):
                break
            if re.match(r'^(?:vol(?:ume)?\.?\s*\d+|issn\b|page\s*\d+)', p, re.I):
                continue
            if any(term in p.lower() for term in ['department of', 'university', 'hospital', 'institute']):
                continue
            if not title_candidate and len(p) >= 10:
                title_candidate = p
                title_para_idx = idx
                break

        # Check paragraph style for explicit title style if not matched
        if not title_candidate:
            for idx, p in enumerate(doc.paragraphs[:15]):
                txt = p.text.strip()
                if not txt or is_article_type_heading(txt):
                    continue
                style_name = (p.style.name or '').lower() if p.style else ''
                if ('title' in style_name or 'heading 1' in style_name) and len(txt) >= 10:
                    title_candidate = txt
                    title_para_idx = idx
                    break

        # Ultimate fallback
        if not title_candidate and paragraphs:
            for idx, p in enumerate(paragraphs[:5]):
                if not is_article_type_heading(p):
                    title_candidate = p
                    title_para_idx = idx
                    break
            if not title_candidate:
                title_candidate = paragraphs[0]

        metadata["article_title"] = title_candidate
        metadata["detected_article_type"] = detected_article_type

        # 5. Extract Abstract Block
        abstract_lines = []
        in_abstract = False
        for p in paragraphs:
            p_lower = p.lower()
            if p_lower.startswith(('abstract', 'summary')):
                in_abstract = True
                # Include content after 'Abstract:' if present on same line
                after_header = re.sub(r'^(abstract|summary)[:\s\-]*', '', p, flags=re.IGNORECASE).strip()
                if after_header:
                    abstract_lines.append(after_header)
                continue

            if in_abstract:
                if p_lower.startswith(('keywords', 'key words', '1. introduction', '1. ', 'introduction')):
                    in_abstract = False
                    break
                abstract_lines.append(p)

        if abstract_lines:
            metadata["abstract"] = " ".join(abstract_lines)

        # 6. Extract Keywords
        for p in paragraphs:
            p_lower = p.lower()
            if p_lower.startswith(('keywords:', 'key words:', 'keywords', 'key words')):
                kw_str = re.sub(r'^(keywords|key words)[:\s\-]*', '', p, flags=re.IGNORECASE).strip()
                if kw_str:
                    raw_kws = re.split(r'[,;•\n]', kw_str)
                    metadata["keywords"] = [k.strip() for k in raw_kws if k.strip()]
                break

        # 7. Extract Authors & Affiliations
        # Scan paragraphs after title_para_idx up to Abstract
        start_idx = (title_para_idx + 1) if title_para_idx >= 0 else 1
        post_title_paras = paragraphs[start_idx:start_idx + 8]
        author_candidates = []
        affiliation_candidates = []

        for p in post_title_paras:
            p_lower = p.lower()
            if p_lower.startswith(('abstract', 'summary', 'keywords', 'doi:')):
                break
            if any(term in p_lower for term in ['university', 'department', 'institute', 'hospital', 'college', 'school', 'center', 'centre', 'laboratory', 'inc.', 'ltd.', 'via pinto', 'italy', 'usa', 'uk']):
                affiliation_candidates.append(p)
            elif '@' not in p and len(p.split()) < 35:
                author_candidates.append(p)

        if author_candidates:
            first_author_line = author_candidates[0]
            authors = [re.sub(r'[\d,\*\†\‡\§]+', '', a).strip() for a in re.split(r'[,;&]|\band\b', first_author_line)]
            authors = [a for a in authors if a and len(a) > 2]
            metadata["lead_author"] = authors[0] if authors else first_author_line
            if len(authors) > 1:
                metadata["co_authors"] = authors[1:]
            else:
                metadata["co_authors"] = author_candidates[1:]

        metadata["affiliations"] = affiliation_candidates

    except Exception as e:
        print(f"Error parsing DOCX metadata: {e}")

    return metadata


def process_uploaded_zip_package(zip_path: str, extract_target_dir: str) -> Dict[str, Any]:
    """
    Extracts ZIP manuscript archive and processes contained DOCX files and art assets.
    Returns extracted metadata and list of unzipped asset files.
    """
    extracted_files: List[Dict[str, str]] = []
    primary_docx_path: Optional[str] = None
    result_metadata: Dict[str, Any] = {}

    if not zipfile.is_zipfile(zip_path):
        raise ValueError("Uploaded file is not a valid ZIP package")

    os.makedirs(extract_target_dir, exist_ok=True)

    with zipfile.ZipFile(zip_path, 'r') as zip_ref:
        zip_ref.extractall(extract_target_dir)

        for root, _, files in os.walk(extract_target_dir):
            for file in files:
                full_path = os.path.join(root, file)
                rel_path = os.path.relpath(full_path, extract_target_dir)
                ext = os.path.splitext(file)[1].lower()

                category = "Manuscript"
                if ext in ['.png', '.jpg', '.jpeg', '.tif', '.tiff', '.eps', '.svg']:
                    category = "Art"
                elif ext in ['.xml', '.jats']:
                    category = "XML"
                elif ext in ['.pdf']:
                    category = "Proof"

                extracted_files.append({
                    "filename": file,
                    "rel_path": rel_path,
                    "full_path": full_path,
                    "category": category,
                    "ext": ext
                })

                if ext == '.docx' and not file.startswith('~$'):
                    if not primary_docx_path or 'manuscript' in file.lower() or 'main' in file.lower():
                        primary_docx_path = full_path

    # Extract metadata from primary DOCX manuscript if found
    if primary_docx_path:
        result_metadata = extract_docx_metadata_advanced(primary_docx_path)

    result_metadata["primary_docx_path"] = primary_docx_path
    result_metadata["extracted_files"] = extracted_files
    return result_metadata
