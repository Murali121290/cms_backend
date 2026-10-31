"""Manuscript (MS) Analysis & Editorial Log (ELOG) module for Journal Production.

Implements Stage 1 of the 8-stage journal workflow:
- 7-step manuscript checklist:
  1. Overview & Frontmatter (words, paragraphs, pages, title, DOI)
  2. Structure & Hierarchy (headings, IMRAD sections)
  3. Mandatory Declarations (Funding, Ethics, Conflicts, Data availability, Consent)
  4. Asset Cross-Verification (figures/tables cited vs. art files/tables present)
  5. Citation & Reference Integrity (in-text citations vs. bibliography)
  6. Advisory Language Editing Level (3-level scale: Light, Medium, Heavy based on density / 1k words)
  7. Consolidated Author Queries (AQ) & Editorial Log (ELOG) report generation
"""
import os
import re
import json
import logging
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional, Tuple

import docx
from sqlalchemy.orm import Session

from app.domains.journals.models import Journal, JournalArticle, JournalFile, JournalStageDetail, JournalClient
from app.domains.journals.service import MS_ANALYSIS, PRE_EDITING, advance_article_stage
from app.domains.journals.files import article_dir, save_version
from app.domains.journals.metadata_extractor import is_article_type_heading

logger = logging.getLogger(__name__)

# Advisory Language Editing thresholds per 1,000 words
LE_LIGHT_MAX = 10.0      # < 10 errors / 1,000 words: Level 1 – Light
LE_MEDIUM_MAX = 25.0     # 10 - 25 errors / 1,000 words: Level 2 – Medium
# > 25 errors / 1,000 words: Level 3 – Heavy


def _translate_path(path: Optional[str]) -> Optional[str]:
    if not path:
        return None
    if os.path.exists(path):
        return path
    # Container -> Host: /opt/cms_runtime/data -> D:/cms_backend/data/cms_runtime_data
    if path.startswith("/opt/cms_runtime/data"):
        win_path = path.replace("/opt/cms_runtime/data", "d:/cms_backend/data/cms_runtime_data").replace("/", os.sep)
        if os.path.exists(win_path):
            return win_path
    # Host -> Container: D:/cms_backend/data/cms_runtime_data -> /opt/cms_runtime/data
    if "cms_runtime_data" in path.replace("\\", "/"):
        sub = path.replace("\\", "/").split("cms_runtime_data")[-1]
        c_path = f"/opt/cms_runtime/data{sub}"
        if os.path.exists(c_path):
            return c_path
    return None


def resolve_manuscript_docx(db: Session, article: JournalArticle) -> Optional[str]:
    """Finds the actual manuscript DOCX file path on disk."""
    p = _translate_path(article.original_docx_path)
    if p:
        return p
    
    # Check journal_files
    manuscript_file = db.query(JournalFile).filter(
        JournalFile.article_id == article.id,
        JournalFile.category.in_(["Manuscript", "DOCX", "original"])
    ).order_by(JournalFile.version.desc()).first()

    if manuscript_file:
        p = _translate_path(manuscript_file.path)
        if p:
            return p

    return None


def inspect_image_file(file_path: str) -> Dict[str, Any]:
    """Inspects an image file to check format, dimensions, color mode, and DPI."""
    try:
        from PIL import Image
        with Image.open(file_path) as img:
            dpi = img.info.get("dpi")
            dpi_val = int(dpi[0]) if dpi and isinstance(dpi, (tuple, list)) and dpi[0] else None
            return {
                "width": img.width,
                "height": img.height,
                "format": img.format,
                "mode": img.mode,
                "dpi": dpi_val,
                "is_low_res": bool(dpi_val and dpi_val < 300)
            }
    except Exception:
        return {}


def calculate_le_density(text: str, total_words: int) -> Tuple[int, float, str, List[Dict[str, Any]]]:
    """
    Computes linguistic errors/issues count and density per 1,000 words.
    Returns: (error_count, density, recommended_level, sample_issues)
    """
    if not text or total_words == 0:
        return 0, 0.0, "Level 1 – Light", []

    issues: List[Dict[str, Any]] = []

    # 1. Punctuation spacing errors (space before comma/period or missing space after)
    space_before_punct = re.finditer(r'\s+([,;:.!?])', text)
    for m in space_before_punct:
        issues.append({
            "type": "punctuation_spacing",
            "message": f"Unwanted space before punctuation '{m.group(1)}'",
            "snippet": text[max(0, m.start() - 15):min(len(text), m.end() + 15)]
        })

    missing_space_after = re.finditer(r'([,;:])(?=[A-Za-z])', text)
    for m in missing_space_after:
        issues.append({
            "type": "missing_space",
            "message": f"Missing space after punctuation '{m.group(1)}'",
            "snippet": text[max(0, m.start() - 15):min(len(text), m.end() + 15)]
        })

    # 2. Multiple consecutive spaces
    multi_spaces = re.finditer(r'[ \t]{2,}', text)
    for m in multi_spaces:
        issues.append({
            "type": "multiple_spaces",
            "message": "Multiple consecutive spaces found",
            "snippet": text[max(0, m.start() - 15):min(len(text), m.end() + 15)]
        })

    # 3. Long rambling sentences (> 40 words)
    sentences = re.split(r'[.!?]+\s+', text)
    for s in sentences:
        words = s.strip().split()
        if len(words) > 40:
            issues.append({
                "type": "long_sentence",
                "message": f"Long sentence ({len(words)} words) may require syntactic restructuring",
                "snippet": " ".join(words[:15]) + "..."
            })

    # 4. Repeated adjacent words (e.g., "the the", "in in")
    repeated_words = re.finditer(r'\b([a-zA-Z]{2,})\s+\1\b', text, re.IGNORECASE)
    for m in repeated_words:
        issues.append({
            "type": "repeated_word",
            "message": f"Repeated word '{m.group(1)}'",
            "snippet": text[max(0, m.start() - 15):min(len(text), m.end() + 15)]
        })

    # 5. Common confusing pairs / passive voice markers
    confused_patterns = [
        (r'\bwhich\b(?=[^,]*\bthat\b)', "Possible restrictive clause misuse ('which' vs 'that')"),
        (r'\bthere\s+is\s+a\s+need\s+to\b', "Wordy passive phrasing ('there is a need to')"),
        (r'\bin\s+order\s+to\b', "Wordy phrasing ('in order to' -> 'to')"),
        (r'\bdue\s+to\s+the\s+fact\s+that\b', "Wordy phrase ('due to the fact that' -> 'because')"),
    ]
    for pattern, msg in confused_patterns:
        for m in re.finditer(pattern, text, re.IGNORECASE):
            issues.append({
                "type": "style_refinement",
                "message": msg,
                "snippet": text[max(0, m.start() - 15):min(len(text), m.end() + 15)]
            })

    total_issues = len(issues)
    density = round((total_issues / max(1, total_words)) * 1000.0, 1)

    if density < LE_LIGHT_MAX:
        recommended = "Level 1 – Light"
    elif density <= LE_MEDIUM_MAX:
        recommended = "Level 2 – Medium"
    else:
        recommended = "Level 3 – Heavy"

    return total_issues, density, recommended, issues[:30]


ARTICLE_TYPE_SPECS: Dict[str, Dict[str, Any]] = {
    "Original Article": {"min_words": 3000, "max_words": 6500, "max_abstract": 120, "target_refs": 50, "allow_art": True},
    "Research Article": {"min_words": 3000, "max_words": 6500, "max_abstract": 120, "target_refs": 50, "allow_art": True},
    "Review Article": {"min_words": 4000, "max_words": 6000, "max_abstract": 120, "target_refs": 80, "allow_art": True, "required_sections": ["Future perspective", "Executive summary"]},
    "Case Study": {"min_words": 1500, "max_words": 3000, "max_abstract": 120, "target_refs": 30, "allow_art": True},
    "Case Report": {"min_words": 1500, "max_words": 3000, "max_abstract": 120, "target_refs": 30, "allow_art": True},
    "Short Communication": {"min_words": 1500, "max_words": 3000, "max_abstract": 120, "target_refs": 30, "allow_art": True},
    "Special Report": {"min_words": 1500, "max_words": 3000, "max_abstract": 120, "target_refs": 50, "allow_art": True},
    "Editorial": {"min_words": 500, "max_words": 1500, "max_abstract": 120, "max_refs": 20, "allow_art": False},
    "Commentary": {"min_words": 1500, "max_words": 3000, "max_refs": 20, "allow_art": False},
    "Letter to Editor": {"min_words": 500, "max_words": 1500, "max_refs": 10, "allow_art": False},
    "Methodology": {"min_words": 3000, "max_words": 6500, "max_abstract": 120, "target_refs": 50, "allow_art": True},
    "Clinical Trial": {"min_words": 3500, "max_words": 7000, "max_abstract": 120, "target_refs": 50, "allow_art": True},
}


def compute_ita_compliance(data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Computes the 100-Point Instructions to Authors (ITA) Compliance Score and Component Weightage
    based on the PA Editorial & STM benchmark guidelines (Weightage_Split.docx & ITA complaiance.docx):
    - Front Matter: 20%
    - Body Text: 40%
    - Figures & Artwork: 15%
    - Tables: 10%
    - References & Citations: 15%
    """
    art_type = (data.get("editorial_screening") or {}).get("article_type", "Original Article")
    spec = ARTICLE_TYPE_SPECS.get(art_type, ARTICLE_TYPE_SPECS["Original Article"])
    
    title = str(data.get("title") or "")
    is_type_flag = is_article_type_heading(title) or (len(title.split()) <= 2 and len(title) <= 10)
    total_words = int(data.get("word_count") or 0)
    abstract_words = int(data.get("abstract_words") or 0)
    keywords_count = int(data.get("keywords_count") or 0)
    lead_author = str(data.get("lead_author") or "")
    
    # 1. Front Matter (Max 20 pts)
    if is_type_flag:
        title_score = 1  # heavily penalize if title is an article type instead of real title
    elif len(title) <= 120:
        title_score = 5
    else:
        title_score = 3
    max_abs = spec.get("max_abstract", 120)
    abstract_score = 5 if 0 < abstract_words <= max_abs else (3 if abstract_words > 0 else 0)
    keywords_score = 5 if keywords_count >= 5 else (3 if keywords_count >= 3 else (1 if keywords_count > 0 else 0))
    authors_score = 5 if lead_author and lead_author != "Unknown" else 2
    front_matter_score = title_score + abstract_score + keywords_score + authors_score
    
    # 2. Body Text (Max 40 pts)
    min_w = spec.get("min_words", 1500)
    max_w = spec.get("max_words", 6500)
    if min_w <= total_words <= max_w:
        body_words_score = 15
    elif (min_w * 0.7) <= total_words <= (max_w * 1.3):
        body_words_score = 10
    else:
        body_words_score = 5
        
    imrad = data.get("imrad_sections") or {}
    imrad_passed = sum(1 for v in imrad.values() if v)
    body_imrad_score = round((imrad_passed / max(1, len(imrad))) * 10)
    
    decls = data.get("declarations") or {}
    decls_passed = sum(1 for d in decls.values() if d.get("present"))
    body_decl_score = decls_passed * 3  # 5 decls * 3 = 15
    body_text_score = body_words_score + body_imrad_score + body_decl_score
    
    # 3. Figures & Artwork (Max 15 pts)
    figs = data.get("figures") or []
    uncalled = data.get("uncalled_art_files") or []
    if not spec.get("allow_art", True):
        figures_score = 15 if len(figs) == 0 else 0
    elif len(figs) == 0:
        figures_score = 15
    else:
        missing_figs = sum(1 for f in figs if f.get("status") == "missing_file")
        low_res_figs = sum(1 for f in figs if f.get("is_low_res"))
        fig_match_score = 5 if missing_figs == 0 else max(0, 5 - (missing_figs * 2))
        fig_res_score = 5 if low_res_figs == 0 else max(0, 5 - (low_res_figs * 2))
        uncalled_score = 5 if len(uncalled) == 0 else max(0, 5 - len(uncalled))
        figures_score = fig_match_score + fig_res_score + uncalled_score
        
    # 4. Tables (Max 10 pts)
    tbls = data.get("tables") or []
    if not spec.get("allow_art", True) and len(tbls) > 0:
        tables_score = 0
    elif len(tbls) == 0:
        tables_score = 10
    else:
        tables_score = 10 if len(tbls) <= 8 else 6
        
    # 5. References & Citations (Max 15 pts)
    refs_count = int(data.get("references_count") or 0)
    citations_count = int(data.get("citations_count") or 0)
    target_refs = spec.get("target_refs", 50)
    max_refs = spec.get("max_refs")
    if max_refs:
        ref_count_score = 5 if refs_count <= max_refs else 2
    elif (target_refs * 0.7) <= refs_count <= (target_refs * 1.3):
        ref_count_score = 5
    elif refs_count > 0:
        ref_count_score = 3
    else:
        ref_count_score = 0
        
    cit_ratio = round(citations_count / max(1, refs_count), 2)
    cit_score = 5 if citations_count > 0 and cit_ratio >= 0.5 else (2 if citations_count > 0 else 0)
    ref_style_score = 5 if refs_count > 0 else 0
    references_score = ref_count_score + cit_score + ref_style_score
    
    total_score = min(100, max(0, front_matter_score + body_text_score + figures_score + tables_score + references_score))
    
    if total_score >= 80:
        verdict = "Accepted for Production"
        badge = "Passed"
    elif total_score >= 60:
        verdict = "Can Accept with Revisions"
        badge = "Revisions Required"
    else:
        verdict = "Technical Return to Author (TRTA)"
        badge = "TRTA Recommended"
        
    return {
        "score": total_score,
        "verdict": verdict,
        "verdict_badge": badge,
        "breakdown": {
            "front_matter": {"score": front_matter_score, "max": 20, "weight": 20},
            "body_text": {"score": body_text_score, "max": 40, "weight": 40},
            "figures": {"score": figures_score, "max": 15, "weight": 15},
            "tables": {"score": tables_score, "max": 10, "weight": 10},
            "references": {"score": references_score, "max": 15, "weight": 15},
        },
        "metrics": {
            "title": title,
            "title_length": len(title),
            "max_title_chars": 120,
            "is_type_flag": is_type_flag,
            "abstract_words": abstract_words,
            "max_abstract_words": max_abs,
            "keywords_count": keywords_count,
            "keywords": data.get("keywords") or [],
            "total_words": total_words,
            "min_words": min_w,
            "max_words": max_w,
            "citations_count": citations_count,
            "references_count": refs_count,
            "citation_reference_ratio": cit_ratio
        }
    }


def analyze_manuscript(db: Session, article: JournalArticle) -> Dict[str, Any]:
    """Runs the full 7-step MS Analysis & ELOG engine on an article."""
    docx_path = resolve_manuscript_docx(db, article)
    journal = db.query(Journal).filter(Journal.id == article.journal_id).first()
    client = db.query(JournalClient).filter(JournalClient.id == journal.client_id).first() if journal else None

    # Default baseline stats if no docx is available
    paragraphs_text: List[str] = []
    paragraph_meta: List[Dict[str, Any]] = []
    headings: List[Dict[str, Any]] = []
    table_count = 0

    if docx_path and os.path.exists(docx_path):
        try:
            doc = docx.Document(docx_path)
            table_count = len(doc.tables)
            for idx, p in enumerate(doc.paragraphs):
                txt = p.text.strip()
                style_name = (p.style.name if p.style else "").lower()
                has_num_pr = False
                try:
                    has_num_pr = bool(p._p.pPr is not None and p._p.pPr.numPr is not None)
                except Exception:
                    pass

                superscripts = []
                try:
                    for r in p.runs:
                        if (r.font.superscript or (r.style and 'superscript' in r.style.name.lower())) and r.text.strip():
                            superscripts.append(r.text.strip())
                except Exception:
                    pass

                if txt:
                    paragraphs_text.append(txt)
                    paragraph_meta.append({
                        "index": idx,
                        "text": txt,
                        "style": style_name,
                        "has_num_pr": has_num_pr,
                        "superscripts": superscripts
                    })
                    if "heading 1" in style_name or "heading 2" in style_name or "title" in style_name:
                        headings.append({"index": idx, "style": p.style.name if p.style else "Heading", "text": txt[:80]})
                    elif len(txt) < 80 and txt.isupper() and len(txt.split()) < 8:
                        headings.append({"index": idx, "style": "Heading (Caps)", "text": txt})
        except Exception as e:
            logger.warning("Could not parse DOCX for MS analysis: %s", e)

    full_text = "\n".join(paragraphs_text)
    total_words = sum(len(p.split()) for p in paragraphs_text) if paragraphs_text else (article.word_count or 0)
    total_paras = len(paragraphs_text)
    estimated_pages = max(1, round(total_words / 450)) if total_words else 1

    # Update article basic word count if missing
    if not article.word_count and total_words:
        article.word_count = total_words

    # 1. Overview & Frontmatter
    detected_type = None
    real_title = None
    title_idx = -1

    for idx, p in enumerate(paragraphs_text[:15]):
        if is_article_type_heading(p):
            if not detected_type:
                detected_type = p
            continue
        if re.match(r'^(?:abstract|keywords?|introduction|doi:|https?://|www\.)\b', p, re.I):
            break
        if re.match(r'^(?:vol(?:ume)?\.?\s*\d+|issn\b|page\s*\d+)', p, re.I):
            continue
        if any(term in p.lower() for term in ['department of', 'university', 'hospital', 'institute']):
            continue
        if not real_title and len(p) >= 10:
            real_title = p
            title_idx = idx
            break

    # Resolve Title: validate against article type heading
    curr_title = (article.article_title or "").strip()
    if not curr_title or is_article_type_heading(curr_title) or len(curr_title) <= 10:
        if real_title:
            title = real_title
            article.article_title = real_title
        elif paragraphs_text:
            title = paragraphs_text[0]
        else:
            title = "Untitled Article"
    else:
        title = curr_title

    # Resolve Lead Author: ensure it is not the title or an article type
    curr_author = (article.lead_author or "").strip()
    if not curr_author or curr_author == title or is_article_type_heading(curr_author) or curr_author == "Unknown":
        start_search = (title_idx + 1) if title_idx >= 0 else 1
        for p in paragraphs_text[start_search:start_search + 5]:
            if re.match(r'^(?:abstract|keywords?|introduction|doi:)\b', p, re.I):
                break
            if any(term in p.lower() for term in ['department of', 'university', 'hospital', 'institute']):
                continue
            authors_clean = [re.sub(r'[\d,\*\†\‡\§]+', '', a).strip() for a in re.split(r'[,;&]|\band\b', p)]
            authors_clean = [a for a in authors_clean if a and len(a) > 2]
            if authors_clean:
                lead_author = authors_clean[0]
                article.lead_author = lead_author
                break
        else:
            lead_author = curr_author or "Unknown"
    else:
        lead_author = curr_author

    # Sync detected article type to editorial screening if applicable
    if detected_type:
        ed_scr = article.extracted_metadata.get("editorial_screening") if article.extracted_metadata else {}
        if not ed_scr:
            ed_scr = {}
        if not ed_scr.get("article_type") or is_article_type_heading(ed_scr.get("article_type", "")):
            normalized_type = "Review Article" if "review" in detected_type.lower() else (
                "Case Report" if "case" in detected_type.lower() else (
                    "Editorial" if "editorial" in detected_type.lower() else (
                        "Short Communication" if "communication" in detected_type.lower() else (
                            "Special Report" if "special" in detected_type.lower() else detected_type
                        )
                    )
                )
            )
            ed_scr["article_type"] = normalized_type
            if not article.extracted_metadata:
                article.extracted_metadata = {}
            article.extracted_metadata["editorial_screening"] = ed_scr
            if not article.complexity_level or is_article_type_heading(article.complexity_level):
                article.complexity_level = normalized_type

    doi = article.article_doi or "Pending Assignment"
    
    # Abstract extraction & word count
    abstract_text = ""
    abstract_words = 0
    in_abstract = False
    for p in paragraphs_text[:35]:
        p_strip = p.strip()
        if re.match(r'^(?:abstract|summary)\b[:\s]*$', p_strip, re.I):
            in_abstract = True
            continue
        m_abs = re.match(r'^(?:abstract|summary)\b[:\s]*(.+)', p_strip, re.I)
        if m_abs:
            abstract_text = m_abs.group(1).strip()
            abstract_words = len(abstract_text.split())
            break
        if in_abstract:
            if re.match(r'^(?:keywords?|key\s+words?|introduction|background)\b[:\s]*', p_strip, re.I):
                break
            abstract_text += (" " if abstract_text else "") + p_strip
            abstract_words = len(abstract_text.split())
            if abstract_words >= 30 and len(paragraphs_text) > 5:
                in_abstract = False

    # Keywords extraction & count
    keywords_list: List[str] = []
    for p in paragraphs_text[:40]:
        m_kw = re.search(r'\b(?:keywords?|key\s+words?)\s*[:\-–]\s*(.+)', p, re.I)
        if m_kw:
            kw_raw = m_kw.group(1).strip()
            for kw in re.split(r'[,;•·/]', kw_raw):
                cleaned_kw = re.sub(r'[\.\(\)\[\]]', '', kw).strip()
                if cleaned_kw and 1 < len(cleaned_kw) < 60:
                    keywords_list.append(cleaned_kw)
            break
    keywords_count = len(keywords_list)

    # 2. Section Structure (IMRAD)
    imrad_sections = {
        "Abstract": bool(re.search(r'\babstract\b', full_text, re.I)),
        "Introduction": bool(re.search(r'\bintroduction\b', full_text, re.I)),
        "Methods": bool(re.search(r'\b(methods|methodology|materials\s+and\s+methods)\b', full_text, re.I)),
        "Results": bool(re.search(r'\bresults\b', full_text, re.I)),
        "Discussion": bool(re.search(r'\bdiscussion\b', full_text, re.I)),
        "Conclusion": bool(re.search(r'\b(conclusion|conclusions)\b', full_text, re.I)),
        "References": bool(re.search(r'\b(references|bibliography|works\s+cited)\b', full_text, re.I)),
    }

    # 3. Mandatory Declarations
    def _extract_declaration(kw_regex: str, label: str) -> Dict[str, Any]:
        found = False
        snippet = ""
        for p in paragraphs_text:
            if re.search(kw_regex, p, re.I):
                found = True
                snippet = p[:250] + ("..." if len(p) > 250 else "")
                break
        return {
            "present": found,
            "label": label,
            "snippet": snippet if found else None,
            "status": "pass" if found else "warning",
            "message": f"{label} is provided." if found else f"Missing mandatory {label} statement."
        }

    declarations = {
        "funding": _extract_declaration(r'\b(funding|grant\s+support|financial\s+support|funded\s+by)\b', "Funding Statement"),
        "ethics": _extract_declaration(r'\b(ethic(?:al)?\s+approval|ethics\s+committee|institutional\s+review\s+board|irb\b|declaration\s+of\s+helsinki)\b', "Ethical Approval"),
        "conflicts": _extract_declaration(r'\b(conflict(?:s)?\s+of\s+interest|competing\s+interest(?:s)?|financial\s+interest)\b', "Conflict of Interest"),
        "data_availability": _extract_declaration(r'\b(data\s+availability|data\s+sharing|raw\s+data|repository)\b', "Data Availability Statement"),
        "consent": _extract_declaration(r'\b(informed\s+consent|patient\s+consent|consent\s+to\s+participate|consent\s+for\s+publication)\b', "Informed Consent"),
    }

    # 4. Asset Cross-Verification (Figures & Tables)
    fig_matches = set(re.findall(r'\b(?:fig(?:ure)?\.?\s*(\d+[a-z]?))\b', full_text, re.I))
    table_matches = set(re.findall(r'\b(?:table\s*(\d+[a-z]?))\b', full_text, re.I))

    # Also scan table cell text for figure and table callouts
    if docx_path and os.path.exists(docx_path):
        try:
            for t_item in doc.tables:
                for row in t_item.rows:
                    for cell in row.cells:
                        c_txt = cell.text
                        for fm in re.findall(r'\b(?:fig(?:ure)?\.?\s*(\d+[a-z]?))\b', c_txt, re.I):
                            fig_matches.add(fm)
                        for tm in re.findall(r'\b(?:table\s*(\d+[a-z]?))\b', c_txt, re.I):
                            table_matches.add(tm)
        except Exception:
            pass

    # Art files attached
    art_files = db.query(JournalFile).filter(
        JournalFile.article_id == article.id,
        JournalFile.category.in_(["Art", "Figure", "image"])
    ).all()

    figures_list: List[Dict[str, Any]] = []
    for fnum in sorted(fig_matches, key=lambda x: (int(re.findall(r'\d+', x)[0]) if re.findall(r'\d+', x) else 0, x)):
        matched_obj = next((f for f in art_files if f"fig{fnum}" in f.filename.lower() or f"figure_{fnum}" in f.filename.lower() or f"figure{fnum}" in f.filename.lower() or (f.figure_number and str(f.figure_number) == str(fnum))), None)
        matched = matched_obj.filename if matched_obj else None
        fig_info: Dict[str, Any] = {
            "label": f"Figure {fnum}",
            "number": fnum,
            "called_in_text": True,
            "matched_file": matched,
            "status": "ok" if matched else "missing_file"
        }
        if matched_obj and matched_obj.path:
            p = _translate_path(matched_obj.path)
            if p and os.path.exists(p):
                fig_info.update(inspect_image_file(p))
        figures_list.append(fig_info)

    # Detect uncalled artwork files
    matched_file_names = {f["matched_file"] for f in figures_list if f.get("matched_file")}
    uncalled_art_files = [f.filename for f in art_files if f.filename not in matched_file_names]

    tables_list: List[Dict[str, Any]] = []
    for tnum in sorted(table_matches, key=lambda x: (int(re.findall(r'\d+', x)[0]) if re.findall(r'\d+', x) else 0, x)):
        tables_list.append({
            "label": f"Table {tnum}",
            "number": tnum,
            "called_in_text": True,
            "status": "ok"
        })

    # 5. Citation & Reference Integrity
    ref_entries: List[str] = []
    in_ref_section = False
    ref_heading_idx = -1

    for p in paragraph_meta:
        txt = p["text"]
        p_style = p["style"]

        is_ref_heading = False
        if p_style in ['refh', 'reference heading', 'reference_heading', 'references heading']:
            is_ref_heading = True
        elif re.search(r'^(?:\d+[\.\s]+|[IVXLCDM]+[\.\s]+)?(references|bibliography|literature\s+cited|works\s+cited)\b[:\s]*$', txt, re.I):
            is_ref_heading = True

        if is_ref_heading:
            in_ref_section = True
            ref_heading_idx = p["index"]
            continue

        if in_ref_section:
            is_exit = False
            if p_style in ['adrf', 'heading 1', 'heading 2', 'heading 3']:
                is_exit = True
            elif re.search(r'^(correspondence|author\s+details|figure\s+\d+|table\s+\d+|appendix|supplementary\s+material|biographical\s+note)\b[:\s]*', txt, re.I):
                is_exit = True

            if is_exit:
                in_ref_section = False
            else:
                has_num_prefix = bool(re.match(r'^\s*(?:\[\d+\]|\d+[\.\)]|\(\d+\)|\d+\t)\s*', txt))
                is_ref_style = ('ref' in p_style or 'biblio' in p_style) and not is_ref_heading
                has_year_or_doi = bool(re.search(r'\b(19|20)\d{2}\b|doi:\s*10\.\d+', txt, re.I))
                
                if p["has_num_pr"] or has_num_prefix or is_ref_style or (has_year_or_doi and len(txt) > 25) or len(txt) > 30:
                    ref_entries.append(txt)

    if not ref_entries:
        for p in paragraph_meta:
            if ('ref' in p["style"] or 'biblio' in p["style"]) and p["style"] not in ['refh', 'reference heading']:
                ref_entries.append(p["text"])

    ref_count = len(ref_entries)

    # B. Identify Citations called in text (prior to reference section)
    citation_numbers: set = set()
    citation_author_years: set = set()

    for p in paragraph_meta:
        if ref_heading_idx != -1 and p["index"] >= ref_heading_idx:
            break

        txt = p["text"]

        for m in re.finditer(r'\[(\d+(?:[–\-,;\s]+\d+)*)\]|(?<!\w)\((\d+(?:[–\-,;\s]+\d+)*)\)(?!\w)', txt):
            is_bracket = bool(m.group(1))
            group = m.group(1) or m.group(2) or ''
            start = m.start()
            preceding = txt[max(0, start - 15):start]

            if not is_bracket and re.search(r'[\d%=\~±\+\-]\s*$', preceding):
                continue

            for chunk in re.split(r'[,;]', group):
                chunk = chunk.strip()
                range_match = re.match(r'^(\d+)[–\-](\d+)$', chunk)
                if range_match:
                    start_n, end_n = int(range_match.group(1)), int(range_match.group(2))
                    if 0 < start_n <= end_n <= 500:
                        if ref_count == 0 or end_n <= ref_count + 5:
                            citation_numbers.update(range(start_n, end_n + 1))
                elif chunk.isdigit():
                    num = int(chunk)
                    if 0 < num <= 500:
                        if is_bracket or ref_count == 0 or num <= ref_count:
                            citation_numbers.add(num)

        for s_txt in p["superscripts"]:
            for m in re.finditer(r'(\d+(?:[–\-,;\s]+\d+)*)', s_txt):
                for chunk in re.split(r'[,;]', m.group(1)):
                    chunk = chunk.strip()
                    range_match = re.match(r'^(\d+)[–\-](\d+)$', chunk)
                    if range_match:
                        start_n, end_n = int(range_match.group(1)), int(range_match.group(2))
                        if 0 < start_n <= end_n <= 500:
                            citation_numbers.update(range(start_n, end_n + 1))
                    elif chunk.isdigit():
                        num = int(chunk)
                        if 0 < num <= 500:
                            citation_numbers.add(num)

        for m in re.finditer(r'(?<!\w)\(([A-Z][a-zA-Z\s.,&–-]+?(?:19|20)\d{2}[a-z]?(?:;\s*[A-Z][a-zA-Z\s.,&–-]+?(?:19|20)\d{2}[a-z]?)*)\)', txt):
            group = m.group(1).strip()
            if re.search(r'\b(19|20)\d{2}\b', group):
                citation_author_years.add(group)

    # B2. Identify Citations called inside all Tables (review tables, study characteristics, etc.)
    if docx_path and os.path.exists(docx_path):
        try:
            for t_obj in doc.tables:
                for row in t_obj.rows:
                    for cell in row.cells:
                        c_txt = cell.text.strip()
                        if not c_txt:
                            continue

                        # Bracketed or parenthetical numbers
                        for m in re.finditer(r'\[(\d+(?:[–\-,;\s]+\d+)*)\]|(?<!\w)\((\d+(?:[–\-,;\s]+\d+)*)\)(?!\w)', c_txt):
                            is_bracket = bool(m.group(1))
                            group = m.group(1) or m.group(2) or ''
                            for chunk in re.split(r'[,;]', group):
                                chunk = chunk.strip()
                                range_match = re.match(r'^(\d+)[–\-](\d+)$', chunk)
                                if range_match:
                                    start_n, end_n = int(range_match.group(1)), int(range_match.group(2))
                                    if 0 < start_n <= end_n <= 500:
                                        if ref_count == 0 or end_n <= ref_count + 5:
                                            citation_numbers.update(range(start_n, end_n + 1))
                                elif chunk.isdigit():
                                    num = int(chunk)
                                    if 0 < num <= 500:
                                        if is_bracket or ref_count == 0 or num <= ref_count:
                                            citation_numbers.add(num)

                        # Author-year citations
                        for m in re.finditer(r'(?<!\w)\(([A-Z][a-zA-Z\s.,&–-]+?(?:19|20)\d{2}[a-z]?(?:;\s*[A-Z][a-zA-Z\s.,&–-]+?(?:19|20)\d{2}[a-z]?)*)\)', c_txt):
                            group = m.group(1).strip()
                            if re.search(r'\b(19|20)\d{2}\b', group):
                                citation_author_years.add(group)

                        # Superscript runs in table cells
                        for p_cell in cell.paragraphs:
                            for r in p_cell.runs:
                                if (r.font.superscript or (r.style and 'superscript' in r.style.name.lower())) and r.text.strip():
                                    for m in re.finditer(r'(\d+(?:[–\-,;\s]+\d+)*)', r.text.strip()):
                                        for chunk in re.split(r'[,;]', m.group(1)):
                                            chunk = chunk.strip()
                                            range_match = re.match(r'^(\d+)[–\-](\d+)$', chunk)
                                            if range_match:
                                                start_n, end_n = int(range_match.group(1)), int(range_match.group(2))
                                                if 0 < start_n <= end_n <= 500:
                                                    citation_numbers.update(range(start_n, end_n + 1))
                                            elif chunk.isdigit():
                                                num = int(chunk)
                                                if 0 < num <= 500:
                                                    citation_numbers.add(num)
        except Exception as e:
            logger.warning("Error scanning tables for citations: %s", e)

    if ref_count == 0 and citation_numbers:
        ref_count = max(citation_numbers)

    citations_count = len(citation_numbers) if citation_numbers else len(citation_author_years)
    cit_ratio = round(citations_count / max(1, ref_count), 2) if ref_count > 0 else 0.0

    # 5b. Editorial Scope, Article Type & Quality Screening
    existing_meta = article.extracted_metadata or {}
    editorial_screening = dict(existing_meta.get("editorial_screening") or {})
    if not editorial_screening:
        editorial_screening = {
            "scope_aligned": True,
            "scope_notes": "Manuscript fits the aims and scope of the journal.",
            "article_type": article.article_type or "Original Article",
            "peer_review_status": "Accepted for Production",
            "similarity_percentage": existing_meta.get("similarity_percentage", 12),
            "similarity_status": "Normal (<20%)",
            "ethics_clearance_confirmed": True,
            "intake_status": "Cleared for Production"
        }

    # 5c. Instructions to Authors (ITA) 100-Point Compliance & Gap Analysis
    art_spec = ARTICLE_TYPE_SPECS.get(editorial_screening.get("article_type", "Original Article"), ARTICLE_TYPE_SPECS["Original Article"])
    analysis_preview = {
        "title": title,
        "word_count": total_words,
        "abstract_words": abstract_words,
        "keywords_count": keywords_count,
        "keywords": keywords_list,
        "lead_author": lead_author,
        "imrad_sections": imrad_sections,
        "declarations": declarations,
        "figures": figures_list,
        "tables": tables_list,
        "uncalled_art_files": uncalled_art_files,
        "references_count": ref_count,
        "citations_count": citations_count,
        "editorial_screening": editorial_screening
    }
    ita_compliance = compute_ita_compliance(analysis_preview)

    # 6. Advisory Language Editing Level Calculation
    err_count, density, recommended_level, sample_issues = calculate_le_density(full_text, total_words)
    
    # Check if manual override exists
    existing_meta = article.extracted_metadata or {}
    chosen_level = existing_meta.get("le_level") or article.complexity_level or recommended_level
    if chosen_level not in ["Level 1 – Light", "Level 2 – Medium", "Level 3 – Heavy"]:
        chosen_level = recommended_level

    # 7. Author Queries (AQ) Generation
    author_queries: List[Dict[str, Any]] = []
    aq_idx = 1

    # Frontmatter and Title queries
    if is_article_type_heading(title) or len(title) <= 10:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Front Matter",
            "query": f"Manuscript title field appears to contain an article classification ('{title}') rather than the complete article title. Please provide the full manuscript title.",
            "status": "Open"
        })
        aq_idx += 1
    elif len(title) > 120:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Front Matter",
            "query": f"Manuscript title is {len(title)} characters long, exceeding the journal limit of 120 characters. Please provide a shortened title or running head.",
            "status": "Open"
        })
        aq_idx += 1

    max_abs_w = art_spec.get("max_abstract", 120)
    if abstract_words > max_abs_w:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Abstract",
            "query": f"Abstract length is {abstract_words} words, which exceeds the journal maximum of {max_abs_w} words. Please condense the abstract to 120 words or fewer.",
            "status": "Open"
        })
        aq_idx += 1

    if keywords_count < 3:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Keywords",
            "query": f"Please provide 5–10 keywords for manuscript indexing and discoverability (currently {keywords_count} detected).",
            "status": "Open"
        })
        aq_idx += 1

    if not declarations["funding"]["present"]:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Declarations",
            "query": "Please provide a formal Funding statement detailing grant numbers and sponsors, or confirm that this research received no external funding.",
            "status": "Open"
        })
        aq_idx += 1

    if not declarations["ethics"]["present"]:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Declarations",
            "query": "Please provide the Ethics Committee or Institutional Review Board (IRB) approval name and reference number, or state if ethics approval was exempt.",
            "status": "Open"
        })
        aq_idx += 1

    if not declarations["conflicts"]["present"]:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Declarations",
            "query": "Please confirm whether any authors have competing financial or non-financial conflicts of interest.",
            "status": "Open"
        })
        aq_idx += 1

    if not declarations["data_availability"]["present"]:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Declarations",
            "query": "Please provide a Data Availability Statement specifying repository URL/accession codes or state that data is available upon request.",
            "status": "Open"
        })
        aq_idx += 1

    for fig in figures_list:
        if fig["status"] == "missing_file":
            author_queries.append({
                "id": f"AQ-{aq_idx:02d}",
                "category": "Artwork",
                "query": f"{fig['label']} is cited in text but no corresponding artwork file was provided in the package. Please supply high-resolution image file (TIFF/EPS/PNG >= 300 DPI).",
                "status": "Open"
            })
            aq_idx += 1
        elif fig.get("is_low_res"):
            dpi_val = fig.get("dpi")
            author_queries.append({
                "id": f"AQ-{aq_idx:02d}",
                "category": "Artwork Resolution",
                "query": f"{fig['label']} ({fig.get('matched_file')}) resolution is {dpi_val} DPI, which is below the publisher 300 DPI print threshold. Please provide high-resolution replacement.",
                "status": "Open"
            })
            aq_idx += 1

    for uncalled in uncalled_art_files:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Artwork",
            "query": f"Artwork file '{uncalled}' was supplied in the package but is not cited anywhere in the manuscript. Please confirm in-text callout placement or remove.",
            "status": "Open"
        })
        aq_idx += 1

    if editorial_screening.get("similarity_percentage", 0) > 25:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Similarity Check",
            "query": f"Plagiarism/similarity check indicates {editorial_screening['similarity_percentage']}% text overlap, which exceeds the standard 20% limit. Please review and clarify cited sources.",
            "status": "Open"
        })
        aq_idx += 1

    if not art_spec.get("allow_art", True) and (len(figures_list) > 0 or len(tables_list) > 0):
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Article Type Fit",
            "query": f"Figures and tables are not permitted in {editorial_screening.get('article_type', 'Editorial')} submissions per journal Instructions to Authors (ITA). Please remove them or request article reclassification.",
            "status": "Open"
        })
        aq_idx += 1

    if chosen_level == "Level 3 – Heavy":
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Language Editing",
            "query": f"Manuscript language density ({density} issues/1k words) indicates Level 3 (Heavy) language editing is advised. Authors will receive edited phrasing for verification.",
            "status": "Open"
        })
        aq_idx += 1

    if ref_count > 0 and citations_count == 0:
        author_queries.append({
            "id": f"AQ-{aq_idx:02d}",
            "category": "Citations & References",
            "query": f"The manuscript contains {ref_count} reference entries in the bibliography, but no corresponding in-text citation callouts were detected in the body text. Please ensure all references are cited in text.",
            "status": "Open"
        })
        aq_idx += 1

    if citation_numbers and ref_count > 0:
        missing_in_biblio = [c for c in citation_numbers if c > ref_count]
        if missing_in_biblio:
            author_queries.append({
                "id": f"AQ-{aq_idx:02d}",
                "category": "Citations & References",
                "query": f"Citation(s) [{', '.join(str(c) for c in sorted(missing_in_biblio)[:5])}] are called in text but exceed the {ref_count} reference entries present in the bibliography. Please supply missing reference details.",
                "status": "Open"
            })
            aq_idx += 1

    # 7-Step Checklist State
    missing_decls = sum(1 for d in declarations.values() if not d["present"])
    missing_figs = sum(1 for f in figures_list if f["status"] == "missing_file")

    checklist = [
        {
            "step": 1,
            "title": "Manuscript Overview & Frontmatter",
            "status": "pass" if (total_words > 0 and title and len(title) <= 120 and keywords_count >= 3) else "warning",
            "details": f"{total_words:,} words · Abstract: {abstract_words}w · {keywords_count} kws · Title: {len(title)} chars"
        },
        {
            "step": 2,
            "title": "Section Structure & Hierarchy",
            "status": "pass" if (imrad_sections["Introduction"] and imrad_sections["Methods"]) else "warning",
            "details": f"{sum(1 for v in imrad_sections.values() if v)} of {len(imrad_sections)} key IMRAD sections identified"
        },
        {
            "step": 3,
            "title": "Mandatory Declarations & Compliance",
            "status": "pass" if missing_decls == 0 else "warning",
            "details": f"{len(declarations) - missing_decls} of {len(declarations)} declarations verified" + (f" ({missing_decls} missing)" if missing_decls else "")
        },
        {
            "step": 4,
            "title": "Asset & Callout Cross-Verification",
            "status": "pass" if missing_figs == 0 else "warning",
            "details": f"{len(figures_list)} figure(s) called · {len(tables_list)} table(s) called" + (f" · {missing_figs} figure art file(s) missing" if missing_figs else "")
        },
        {
            "step": 5,
            "title": "Citation & Reference Integrity",
            "status": "pass" if (citations_count > 0 and ref_count > 0) else "warning",
            "details": f"{citations_count} citations called · {ref_count} references listed (Ratio: {cit_ratio})"
        },
        {
            "step": 6,
            "title": "Advisory Language Editing (LE) Level",
            "status": "pass",
            "details": f"{chosen_level} (Score: {density}/1k words; {err_count} detected issues)"
        },
        {
            "step": 7,
            "title": "Author Queries & Editorial Log (ELOG)",
            "status": "pass",
            "details": f"{len(author_queries)} Author Query item(s) logged · ELOG report ready"
        }
    ]

    analysis_data = {
        "article_id": article.id,
        "title": title,
        "doi": doi,
        "lead_author": lead_author,
        "journal_code": journal.journal_code if journal else "Unknown",
        "journal_title": journal.journal_title if journal else "Unknown",
        "client_code": client.client_code if client else "Unknown",
        "publisher_name": client.publisher_name if client else "Unknown",
        "word_count": total_words,
        "estimated_pages": estimated_pages,
        "paragraphs_count": total_paras,
        "table_count": table_count,
        "abstract_words": abstract_words,
        "abstract_text": abstract_text,
        "keywords": keywords_list,
        "keywords_count": keywords_count,
        "title_length": len(title),
        "citation_reference_ratio": cit_ratio,
        "ita_compliance": ita_compliance,
        "imrad_sections": imrad_sections,
        "headings": headings[:20],
        "declarations": declarations,
        "figures": figures_list,
        "tables": tables_list,
        "uncalled_art_files": uncalled_art_files,
        "citations_count": citations_count,
        "references_count": ref_count,
        "le_error_count": err_count,
        "le_density": density,
        "le_recommended_level": recommended_level,
        "le_chosen_level": chosen_level,
        "editorial_screening": editorial_screening,
        "sample_language_issues": sample_issues,
        "author_queries": author_queries,
        "checklist": checklist,
        "analyzed_at": datetime.now(timezone.utc).isoformat()
    }

    # Generate ELOG HTML
    analysis_data["elog_html"] = generate_elog_html(article, analysis_data)

    # Persist in article metadata
    meta = dict(article.extracted_metadata or {})
    meta["ms_analysis"] = analysis_data
    meta["le_level"] = chosen_level
    article.extracted_metadata = meta
    article.complexity_level = chosen_level
    db.commit()

    # Save ELOG file
    _save_elog_file(db, article, analysis_data["elog_html"])

    return analysis_data


def generate_elog_html(article: JournalArticle, data: Dict[str, Any]) -> str:
    """Generates a professional, print-ready HTML Editorial Log (ELOG) document."""
    queries_rows = ""
    for q in data.get("author_queries", []):
        queries_rows += f"""
        <tr style="border-bottom: 1px solid #e2e8f0;">
          <td style="padding: 10px 12px; font-weight: 700; color: #0284c7; font-family: monospace;">{q['id']}</td>
          <td style="padding: 10px 12px; font-size: 13px; color: #475569;"><span style="background: #f1f5f9; padding: 2px 6px; border-radius: 4px; font-size: 11px;">{q['category']}</span></td>
          <td style="padding: 10px 12px; font-size: 13px; color: #1e293b;">{q['query']}</td>
          <td style="padding: 10px 12px; font-size: 12px; color: #64748b; min-width: 140px; border-left: 1px dashed #cbd5e1;">[ Author response here ]</td>
        </tr>
        """

    decls_rows = ""
    for k, d in data.get("declarations", {}).items():
        status_color = "#16a34a" if d["present"] else "#dc2626"
        status_badge = "CONFIRMED" if d["present"] else "MISSING"
        decls_rows += f"""
        <tr style="border-bottom: 1px solid #e2e8f0;">
          <td style="padding: 8px 12px; font-weight: 600; color: #334155;">{d['label']}</td>
          <td style="padding: 8px 12px;"><span style="color: {status_color}; font-weight: 700; font-size: 12px;">● {status_badge}</span></td>
          <td style="padding: 8px 12px; font-size: 12px; color: #64748b;">{d['snippet'] or 'None provided in manuscript.'}</td>
        </tr>
        """

    figures_rows = ""
    for f in data.get("figures", []):
        if f["status"] == "ok":
            res_txt = f"{f.get('dpi')} DPI" if f.get("dpi") else "Standard"
            if f.get("is_low_res"):
                status_badge = f"<span style='color: #d97706; font-weight: 700;'>Low Res ({res_txt})</span>"
            else:
                status_badge = f"<span style='color: #16a34a; font-weight: 700;'>Matched ({res_txt})</span>"
        else:
            status_badge = "<span style='color: #dc2626; font-weight: 700;'>Missing Art File</span>"

        dims = f"{f.get('width')}x{f.get('height')} px · {f.get('format', '')}" if f.get("width") else "Word embedded / link"
        figures_rows += f"""
        <tr style="border-bottom: 1px solid #e2e8f0;">
          <td style="padding: 6px 12px; font-weight: 600;">{f['label']}</td>
          <td style="padding: 6px 12px; font-size: 12px;">{status_badge}</td>
          <td style="padding: 6px 12px; font-size: 12px; color: #64748b; font-family: monospace;">{f.get('matched_file') or 'No file found in package'} ({dims})</td>
        </tr>
        """

    checklist_cards = ""
    for c in data.get("checklist", []):
        icon = "✓" if c["status"] == "pass" else "⚠"
        color = "#16a34a" if c["status"] == "pass" else "#d97706"
        checklist_cards += f"""
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 10px 14px; margin-bottom: 8px;">
          <div style="display: flex; justify-content: space-between; align-items: center;">
            <div style="font-weight: 600; font-size: 13px; color: #1e293b;">
              <span style="color: {color}; margin-right: 6px; font-weight: 800;">{icon}</span> Step {c['step']}: {c['title']}
            </div>
            <div style="font-size: 12px; color: #64748b;">{c['details']}</div>
          </div>
        </div>
        """

    ed = data.get("editorial_screening") or {}
    scope_badge = "<span style='color: #16a34a; font-weight: 700;'>✓ Confirmed / Aligned</span>" if ed.get("scope_aligned", True) else "<span style='color: #dc2626; font-weight: 700;'>⚠ Scope Query Flagged</span>"
    sim_pct = ed.get("similarity_percentage", 12)
    sim_color = "#16a34a" if sim_pct < 20 else ("#d97706" if sim_pct <= 30 else "#dc2626")

    ita = data.get("ita_compliance") or {}
    ita_score = ita.get("score", 80)
    ita_verdict = ita.get("verdict", "Accepted for Production")
    ita_score_color = "#16a34a" if ita_score >= 80 else ("#d97706" if ita_score >= 60 else "#dc2626")
    breakdown = ita.get("breakdown") or {}
    front_s = breakdown.get("front_matter", {}).get("score", 16)
    body_s = breakdown.get("body_text", {}).get("score", 32)
    fig_s = breakdown.get("figures", {}).get("score", 13)
    tbl_s = breakdown.get("tables", {}).get("score", 8)
    ref_s = breakdown.get("references", {}).get("score", 13)
    cit_ratio = data.get("citation_reference_ratio", 1.0)

    ita_compliance_card = f"""
    <div style="background: #ffffff; border: 2px solid #0284c7; border-radius: 10px; padding: 18px 22px; margin-bottom: 24px; box-shadow: 0 4px 12px rgba(2,132,199,0.08);">
      <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 12px;">
        <div>
          <span style="font-size: 14px; font-weight: 800; color: #0f172a; text-transform: uppercase; letter-spacing: 0.5px;">Instructions to Authors (ITA) 100-Point Compliance &amp; Gap Analysis</span>
          <div style="font-size: 11px; color: #64748b; margin-top: 2px;">PA Editorial Ltd. Weighted Preflight Matrix (Weightage_Split &amp; STM Standard)</div>
        </div>
        <div style="text-align: right;">
          <span style="font-size: 26px; font-weight: 900; color: {ita_score_color};">{ita_score}%</span>
          <div style="font-size: 11px; font-weight: 700; color: {ita_score_color}; text-transform: uppercase;">{ita_verdict}</div>
        </div>
      </div>
      <div style="background: #e2e8f0; border-radius: 9999px; height: 12px; overflow: hidden; margin-bottom: 16px;">
        <div style="background: {ita_score_color}; height: 100%; width: {ita_score}%; border-radius: 9999px;"></div>
      </div>
      <div style="display: grid; grid-template-columns: repeat(5, 1fr); gap: 10px; font-size: 11px; text-align: center;">
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 10px 8px;">
          <div style="color: #64748b; font-weight: 700; text-transform: uppercase; font-size: 10px;">Front Matter (20%)</div>
          <div style="font-size: 15px; font-weight: 800; color: #0f172a; margin: 3px 0;">{front_s} / 20</div>
          <div style="color: #64748b; font-size: 10px;">Title &le;120c, Abs, Kws</div>
        </div>
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 10px 8px;">
          <div style="color: #64748b; font-weight: 700; text-transform: uppercase; font-size: 10px;">Body Text (40%)</div>
          <div style="font-size: 15px; font-weight: 800; color: #0f172a; margin: 3px 0;">{body_s} / 40</div>
          <div style="color: #64748b; font-size: 10px;">Words, IMRAD, 5 Decls</div>
        </div>
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 10px 8px;">
          <div style="color: #64748b; font-weight: 700; text-transform: uppercase; font-size: 10px;">Figures &amp; Art (15%)</div>
          <div style="font-size: 15px; font-weight: 800; color: #0f172a; margin: 3px 0;">{fig_s} / 15</div>
          <div style="color: #64748b; font-size: 10px;">&ge;300 DPI, Calls</div>
        </div>
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 10px 8px;">
          <div style="color: #64748b; font-weight: 700; text-transform: uppercase; font-size: 10px;">Tables (10%)</div>
          <div style="font-size: 15px; font-weight: 800; color: #0f172a; margin: 3px 0;">{tbl_s} / 10</div>
          <div style="color: #64748b; font-size: 10px;">Numbered &amp; Cited</div>
        </div>
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 10px 8px;">
          <div style="color: #64748b; font-weight: 700; text-transform: uppercase; font-size: 10px;">References (15%)</div>
          <div style="font-size: 15px; font-weight: 800; color: #0f172a; margin: 3px 0;">{ref_s} / 15</div>
          <div style="color: #64748b; font-size: 10px;">Ratio: {cit_ratio}</div>
        </div>
      </div>
    </div>
    """

    editorial_screening_card = f"""
    <div style="background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 8px; padding: 14px 18px; margin-bottom: 24px;">
      <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 10px;">
        <span style="font-size: 13px; font-weight: 800; color: #0f172a; text-transform: uppercase; letter-spacing: 0.5px;">Editorial Scope &amp; Quality Screening</span>
        <span style="font-size: 11px; font-weight: 700; background: #e2e8f0; color: #334155; padding: 2px 8px; border-radius: 4px; text-transform: uppercase;">{ed.get('intake_status', 'Cleared for Production')}</span>
      </div>
      <div style="display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; font-size: 12px; color: #334155;">
        <div><strong>Article Type:</strong> {ed.get('article_type', 'Original Article')}</div>
        <div><strong>Aims &amp; Scope:</strong> {scope_badge}</div>
        <div><strong>Similarity Check:</strong> <span style="color: {sim_color}; font-weight: 800;">{sim_pct}%</span> ({ed.get('similarity_status', 'Normal')})</div>
        <div><strong>Peer Review:</strong> {ed.get('peer_review_status', 'Accepted')}</div>
      </div>
      <div style="margin-top: 8px; font-size: 11px; color: #64748b;">
        <strong>Screening Note:</strong> {ed.get('scope_notes', 'All frontmatter and editorial scope checks verified.')}
      </div>
    </div>
    """

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>Editorial Log (ELOG) - {data.get('doi')}</title>
  <style>
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      margin: 0; padding: 40px; background: #ffffff; color: #0f172a; line-height: 1.5;
    }}
    .header {{
      display: flex; justify-content: space-between; align-items: flex-start;
      border-bottom: 3px solid #0284c7; padding-bottom: 20px; margin-bottom: 30px;
    }}
    .brand {{ font-size: 24px; font-weight: 800; color: #0f172a; letter-spacing: -0.5px; }}
    .subtitle {{ font-size: 13px; color: #0284c7; font-weight: 700; text-transform: uppercase; letter-spacing: 1px; }}
    .meta-grid {{
      display: grid; grid-template-columns: repeat(6, 1fr); gap: 12px; margin-bottom: 28px;
    }}
    .meta-card {{
      background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 12px 14px;
    }}
    .meta-label {{ font-size: 11px; text-transform: uppercase; color: #64748b; font-weight: 700; margin-bottom: 4px; }}
    .meta-val {{ font-size: 14px; font-weight: 700; color: #0f172a; }}
    .section-title {{
      font-size: 16px; font-weight: 700; color: #0f172a; margin: 24px 0 12px 0; border-bottom: 1px solid #e2e8f0; padding-bottom: 6px;
    }}
    table {{ width: 100%; border-collapse: collapse; margin-bottom: 20px; }}
    th {{ background: #f1f5f9; text-align: left; padding: 10px 12px; font-size: 12px; text-transform: uppercase; color: #475569; letter-spacing: 0.5px; }}
    .badge {{ display: inline-block; padding: 3px 10px; border-radius: 12px; font-size: 12px; font-weight: 700; }}
    .badge-level {{ background: #e0f2fe; color: #0369a1; border: 1px solid #bae6fd; }}
  </style>
</head>
<body>
  <div class="header">
    <div>
      <div class="brand">PA Editorial Ltd. · Editorial Log (ELOG)</div>
      <div class="subtitle">Stage 1 Manuscript Analysis & Compliance Dossier</div>
    </div>
    <div style="text-align: right;">
      <div style="font-size: 13px; font-weight: 600;">{data.get('client_code')} / {data.get('journal_code')}</div>
      <div style="font-size: 12px; color: #64748b;">Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}</div>
    </div>
  </div>

  <div style="margin-bottom: 24px;">
    <h1 style="font-size: 20px; margin: 0 0 8px 0; color: #0f172a;">{data.get('title')}</h1>
    <div style="font-size: 13px; color: #475569;">
      <strong>Lead Author:</strong> {data.get('lead_author')} &nbsp;|&nbsp;
      <strong>DOI:</strong> {data.get('doi')} &nbsp;|&nbsp;
      <strong>Publisher:</strong> {data.get('publisher_name')}
    </div>
  </div>

  <div class="meta-grid">
    <div class="meta-card">
      <div class="meta-label">Total Words</div>
      <div class="meta-val">{data.get('word_count'):,}</div>
    </div>
    <div class="meta-card">
      <div class="meta-label">Abstract / Title</div>
      <div class="meta-val">{data.get('abstract_words', 0)}w abs / {data.get('title_length', 0)}c</div>
    </div>
    <div class="meta-card">
      <div class="meta-label">ITA Score</div>
      <div class="meta-val"><span class="badge" style="background: {ita_score_color}15; color: {ita_score_color}; border: 1px solid {ita_score_color}40;">{ita_score}%</span></div>
    </div>
    <div class="meta-card">
      <div class="meta-label">Advisory LE Level</div>
      <div class="meta-val"><span class="badge badge-level">{data.get('le_chosen_level')}</span></div>
    </div>
    <div class="meta-card">
      <div class="meta-label">Tables &amp; Figs</div>
      <div class="meta-val">{len(data.get('tables', []))} tbl / {len(data.get('figures', []))} fig</div>
    </div>
    <div class="meta-card">
      <div class="meta-label">Refs &amp; Citations</div>
      <div class="meta-val">{data.get('references_count', 0)} refs (Ratio: {cit_ratio})</div>
    </div>
  </div>

  {ita_compliance_card}

  {editorial_screening_card}

  <div class="section-title">1. 7-Step Manuscript Analysis Checklist</div>
  <div style="margin-bottom: 20px;">
    {checklist_cards}
  </div>

  <div class="section-title">2. Mandatory Declarations & Ethics Compliance</div>
  <table>
    <thead>
      <tr>
        <th style="width: 28%;">Declaration Type</th>
        <th style="width: 18%;">Status</th>
        <th style="width: 54%;">Detected Manuscript Statement</th>
      </tr>
    </thead>
    <tbody>
      {decls_rows}
    </tbody>
  </table>

  <div class="section-title">3. Asset Cross-Verification (Artwork & Tables)</div>
  <table>
    <thead>
      <tr>
        <th style="width: 28%;">Asset</th>
        <th style="width: 25%;">Package Status</th>
        <th style="width: 47%;">Matched Artwork File</th>
      </tr>
    </thead>
    <tbody>
      {figures_rows or '<tr><td colspan="3" style="padding: 10px; color: #64748b;">No figures called in manuscript text.</td></tr>'}
    </tbody>
  </table>

  <div class="section-title">4. Consolidated Author Queries (AQ) ({len(data.get('author_queries', []))} items)</div>
  <table>
    <thead>
      <tr>
        <th style="width: 10%;">Query ID</th>
        <th style="width: 16%;">Category</th>
        <th style="width: 48%;">Query Text</th>
        <th style="width: 26%;">Author Response</th>
      </tr>
    </thead>
    <tbody>
      {queries_rows or '<tr><td colspan="4" style="padding: 10px; color: #16a34a; font-weight: 600;">No author queries required. All compliance and asset checks passed!</td></tr>'}
    </tbody>
  </table>

  <div style="margin-top: 40px; padding-top: 15px; border-top: 1px solid #cbd5e1; font-size: 11px; color: #94a3b8; text-align: center;">
    PA Editorial Ltd. · NinjaInkflow Publishing Platform · Manuscript Analysis & Editorial Log Standard
  </div>
</body>
</html>
"""
    return html


def _save_elog_file(db: Session, article: JournalArticle, html_content: str) -> None:
    """Saves the ELOG HTML report into the article working directory as a JournalFile."""
    try:
        slug = f"article_{article.id}_ELOG.html"
        save_version(db, article, "ELOG", slug, html_content.encode("utf-8"), "elog")
        db.commit()
    except Exception as e:
        logger.warning("Could not persist ELOG file to disk: %s", e)


def get_ms_analysis(db: Session, article: JournalArticle) -> Dict[str, Any]:
    """Retrieves cached MS Analysis data or runs analysis on demand."""
    meta = article.extracted_metadata or {}
    analysis = meta.get("ms_analysis")
    if analysis and isinstance(analysis, dict):
        return analysis
    return analyze_manuscript(db, article)


def set_le_level(db: Session, article: JournalArticle, level: str, reason: Optional[str] = None) -> Dict[str, Any]:
    """Sets the Language Editing (LE) Level manually (Light, Medium, Heavy)."""
    valid = ["Level 1 – Light", "Level 2 – Medium", "Level 3 – Heavy"]
    if level not in valid:
        # Match fuzzy
        if "light" in level.lower() or "1" in level:
            level = "Level 1 – Light"
        elif "heavy" in level.lower() or "3" in level:
            level = "Level 3 – Heavy"
        else:
            level = "Level 2 – Medium"

    meta = dict(article.extracted_metadata or {})
    meta["le_level"] = level
    if reason:
        meta["le_level_reason"] = reason

    if "ms_analysis" in meta and isinstance(meta["ms_analysis"], dict):
        meta["ms_analysis"]["le_chosen_level"] = level
        # Update checklist step 6
        for c in meta["ms_analysis"].get("checklist", []):
            if c.get("step") == 6:
                density = meta["ms_analysis"].get("le_density", 0.0)
                err_count = meta["ms_analysis"].get("le_error_count", 0)
                c["details"] = f"{level} (Score: {density}/1k words; {err_count} detected issues)"
        meta["ms_analysis"]["elog_html"] = generate_elog_html(article, meta["ms_analysis"])
        _save_elog_file(db, article, meta["ms_analysis"]["elog_html"])

    article.extracted_metadata = meta
    article.complexity_level = level
    db.commit()

    return {
        "status": "success",
        "article_id": article.id,
        "le_level": level,
        "reason": reason
    }


def complete_ms_analysis_stage(db: Session, article: JournalArticle, user_id: Optional[int] = None, remarks: Optional[str] = None) -> Dict[str, Any]:
    """
    Completes Stage 1 (MS Analysis & ELOG) and advances to Stage 2 (Pre-Editing).
    Ensures analysis and ELOG are generated before advancing.
    """
    analysis = get_ms_analysis(db, article)
    adv_remarks = remarks or f"MS Analysis & ELOG completed. LE Level: {analysis.get('le_chosen_level')}, {len(analysis.get('author_queries', []))} AQs logged."
    
    res = advance_article_stage(db, article.id, adv_remarks)
    return {
        **res,
        "ms_analysis": analysis
    }


def set_editorial_screening(db: Session, article: JournalArticle, screening_data: Dict[str, Any]) -> Dict[str, Any]:
    """Updates the editorial screening criteria (scope, article type, similarity, clearance) and refreshes ELOG."""
    meta = dict(article.extracted_metadata or {})
    ms = dict(meta.get("ms_analysis") or {})
    if not ms:
        ms = analyze_manuscript(db, article)
        meta = dict(article.extracted_metadata or {})
        ms = dict(meta.get("ms_analysis") or {})

    current_screening = dict(ms.get("editorial_screening") or {})
    current_screening.update(screening_data)

    if "similarity_percentage" in screening_data:
        pct = int(screening_data["similarity_percentage"])
        current_screening["similarity_percentage"] = pct
        current_screening["similarity_status"] = "Normal (<20%)" if pct < 20 else ("Moderate (20-30%)" if pct <= 30 else "High (>30%)")
        meta["similarity_percentage"] = pct

    if "article_type" in screening_data:
        article.article_type = str(screening_data["article_type"])
        current_screening["article_type"] = article.article_type

    ms["editorial_screening"] = current_screening
    ms["ita_compliance"] = compute_ita_compliance(ms)

    for c in ms.get("checklist", []):
        if c.get("step") == 2:
            imrad_sections = ms.get("imrad_sections", {})
            c["details"] = f"Type: {current_screening.get('article_type', 'Original Article')} · {sum(1 for v in imrad_sections.values() if v)}/{len(imrad_sections)} IMRAD sections"
            c["status"] = "pass" if current_screening.get("scope_aligned", True) else "warning"

    ms["elog_html"] = generate_elog_html(article, ms)
    meta["ms_analysis"] = ms
    article.extracted_metadata = meta
    db.commit()

    _save_elog_file(db, article, ms["elog_html"])
    return ms


def trigger_trta_package(db: Session, article: JournalArticle, reason: Optional[str] = None) -> Dict[str, Any]:
    """Records Technical Return to Author (TRTA) and compiles author return package metadata."""
    meta = dict(article.extracted_metadata or {})
    ms = dict(meta.get("ms_analysis") or {})
    if not ms:
        ms = analyze_manuscript(db, article)
        meta = dict(article.extracted_metadata or {})
        ms = dict(meta.get("ms_analysis") or {})

    screening = dict(ms.get("editorial_screening") or {})
    screening["intake_status"] = "Technical Return to Author (TRTA)"
    screening["trta_reason"] = reason or "Technical or mandatory declaration deficiencies detected during intake pre-flight."
    screening["trta_timestamp"] = datetime.now(timezone.utc).isoformat()

    ms["editorial_screening"] = screening
    meta["ms_analysis"] = ms
    meta["intake_decision"] = "TRTA"
    article.extracted_metadata = meta
    db.commit()

    return {
        "status": "TRTA",
        "message": f"Article #{article.id} marked as Technical Return to Author (TRTA).",
        "author_queries_count": len(ms.get("author_queries", [])),
        "missing_declarations": sum(1 for d in ms.get("declarations", {}).values() if not d.get("present")),
        "missing_figures": sum(1 for f in ms.get("figures", []) if f.get("status") == "missing_file"),
        "trta_timestamp": screening["trta_timestamp"]
    }
