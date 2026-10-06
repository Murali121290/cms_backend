#!/usr/bin/env python3
"""
Usage:
    python manuscript_to_jats.py input.docx -p jmir_mededu_profile.json -o output.xml --dtd path/to/journalpublishing.dtd --warnings warnings.txt
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import subprocess
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator, Optional, Sequence, Union

from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from lxml import etree

XLINK = "http://www.w3.org/1999/xlink"
MML = "http://www.w3.org/1998/Math/MathML"
XML_NS = "http://www.w3.org/XML/1998/namespace"
perl_script_path = "utf8.pl"

MONTHS = {
    "jan": "01", "january": "01", "feb": "02", "february": "02",
    "mar": "03", "march": "03", "apr": "04", "april": "04",
    "may": "05", "jun": "06", "june": "06", "jul": "07", "july": "07",
    "aug": "08", "august": "08", "sep": "09", "sept": "09", "september": "09",
    "oct": "10", "october": "10", "nov": "11", "november": "11",
    "dec": "12", "december": "12",
}


# ----------------------------- utilities -----------------------------

def local_name(tag: str) -> str:
    return etree.QName(tag).localname


def normalize_space(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def style_name(p: Paragraph) -> str:
    return p.style.name if p.style else ""


def iter_block_items(parent: Union[_Document, _Cell]) -> Iterator[Union[Paragraph, Table]]:
    """Yield Paragraph and Table objects in true document order."""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
        parent_obj = parent
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
        parent_obj = parent
    else:
        raise TypeError(type(parent))

    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent_obj)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent_obj)


def append_text(parent: etree._Element, text: str) -> None:
    """Append text after current last child, preserving mixed content."""
    if not text:
        return
    if len(parent):
        last = parent[-1]
        last.tail = (last.tail or "") + text
    else:
        parent.text = (parent.text or "") + text


def wrap_text(parent: etree._Element, tag: str, text: str) -> etree._Element:
    el = etree.SubElement(parent, tag)
    el.text = text
    return el


def date_parts(s: str) -> Optional[tuple[str, str, str]]:
    """Parse 23.Mar.2026, 28.Sep.2026, Jul 23, 2024, etc."""
    s = normalize_space(s).strip(".;")
    m = re.search(r"(?i)\b(\d{1,2})[.\-/ ]([A-Za-z]{3,9}|\d{1,2})[.\-/ ](\d{4})\b", s)
    if m:
        day, mon, year = m.groups()
        month = MONTHS.get(mon.lower(), mon.zfill(2))
        return day.zfill(2), month, year
    m = re.search(r"(?i)\b([A-Za-z]{3,9})\s+(\d{1,2}),\s*(\d{4})\b", s)
    if m:
        mon, day, year = m.groups()
        return day.zfill(2), MONTHS.get(mon.lower(), ""), year
    return None


def add_date(parent: etree._Element, date_type: Optional[str], parts: tuple[str, str, str]) -> etree._Element:
    attrs = {"date-type": date_type} if date_type else {}
    el = etree.SubElement(parent, "date", **attrs)
    d, m, y = parts
    etree.SubElement(el, "day").text = str(int(d)) if d else ""
    etree.SubElement(el, "month").text = m
    etree.SubElement(el, "year").text = y
    return el


def person_name(parent: etree._Element, full_name: str, role: Optional[str] = None) -> etree._Element:
    """Heuristic Western name parser: surname = last token; given names = preceding tokens."""
    full_name = normalize_space(full_name)
    bits = full_name.split()
    if not bits:
        return etree.SubElement(parent, "name", {"name-style": "western"})
    surname = bits[-1]
    given = " ".join(bits[:-1])
    name = etree.SubElement(parent, "name", {"name-style": "western"})
    etree.SubElement(name, "surname").text = surname
    if given:
        etree.SubElement(name, "given-names").text = given
    return name


def citation_name(parent: etree._Element, raw: str) -> None:
    raw = normalize_space(raw)
    # given initials are conventionally the final token(s): "van der Molen HT", "Artino AR Jr"
    bits = raw.split()
    if not bits:
        return
    suffix = ""
    if bits[-1] in {"Jr", "Sr", "II", "III", "IV"} and len(bits) >= 3:
        suffix = " " + bits.pop()
    given = bits.pop() if bits else ""
    surname = " ".join(bits)
    name = etree.SubElement(parent, "name", {"name-style": "western"})
    etree.SubElement(name, "surname").text = surname
    etree.SubElement(name, "given-names").text = given + suffix


def add_person_group(parent: etree._Element, raw: str, group_type: str) -> etree._Element:
    pg = etree.SubElement(parent, "person-group", {"person-group-type": group_type})
    raw = normalize_space(raw)
    has_etal = bool(re.search(r"\bet al\.?$", raw, flags=re.I))
    raw = re.sub(r",?\s*et al\.?$", "", raw, flags=re.I)
    for item in [x.strip() for x in raw.split(",") if x.strip()]:
        citation_name(pg, item)
    if has_etal:
        etree.SubElement(pg, "etal")
    return pg


def set_xlink(el: etree._Element, attr: str, value: str) -> None:
    el.set(f"{{{XLINK}}}{attr}", value)


# -------------------------- Word inline model --------------------------

@dataclass
class Span:
    text: str
    bold: bool = False
    italic: bool = False
    sup: bool = False
    sub: bool = False
    href: Optional[str] = None


def paragraph_spans(p: Paragraph) -> list[Span]:
    """Read w:r and w:hyperlink content in order, including formatting."""
    spans: list[Span] = []
    rels = p.part.rels
    for child in p._p.iterchildren():
        tag = local_name(child.tag)
        href = None
        runs = []
        if tag == "r":
            runs = [child]
        elif tag == "hyperlink":
            rid = child.get(qn("r:id"))
            if rid and rid in rels:
                href = rels[rid].target_ref
            runs = [r for r in child if local_name(r.tag) == "r"]
        else:
            continue
        for r in runs:
            texts = []
            for node in r.iterchildren():
                n = local_name(node.tag)
                if n in {"t", "delText", "instrText"}:
                    texts.append(node.text or "")
                elif n == "tab":
                    texts.append("\t")
                elif n in {"br", "cr"}:
                    texts.append("\n")
            text = "".join(texts)
            if not text:
                continue
            rpr = r.find(qn("w:rPr"))
            bold = italic = sup = sub = False
            if rpr is not None:
                bold = rpr.find(qn("w:b")) is not None
                italic = rpr.find(qn("w:i")) is not None
                va = rpr.find(qn("w:vertAlign"))
                if va is not None:
                    val = va.get(qn("w:val"))
                    sup = val == "superscript"
                    sub = val == "subscript"
            spans.append(Span(text, bold, italic, sup, sub, href))
    if not spans and p.text:
        spans = [Span(p.text)]
    # Word often splits citations such as "[1-4]" into several identical runs.
    # Coalesce adjacent runs with the same semantics so cross-reference parsing
    # operates on the logical text rather than arbitrary DOCX run boundaries.
    merged: list[Span] = []
    for s in spans:
        if merged and (merged[-1].bold, merged[-1].italic, merged[-1].sup, merged[-1].sub, merged[-1].href) == (s.bold, s.italic, s.sup, s.sub, s.href):
            merged[-1].text += s.text
        else:
            merged.append(s)
    return merged


XREF_PATTERN = re.compile(
    r"(\[(?:\d+\s*(?:[-,]\s*\d+)*?)\]|\bFigure\s+\d+\b|\bTable\s+\d+\b|\bMultimedia Appendix\s+\d+\b|\bChecklist\s+\d+\b)",
    flags=re.I,
)


SUPP_FILE_LINE = re.compile(r"^\[[^\]]*\bFile\b[^\]]*\]$", flags=re.I)


def append_xref_token(parent: etree._Element, token: str) -> bool:
    """Append token with JATS xref elements; return True if token recognized."""
    if token.startswith("[") and token.endswith("]"):
        inside = token[1:-1]
        append_text(parent, "[")
        cursor = 0
        for m in re.finditer(r"\d+", inside):
            append_text(parent, inside[cursor:m.start()])
            num = m.group()
            x = etree.SubElement(parent, "xref", {"ref-type": "bibr", "rid": f"ref{num}"})
            x.text = num
            cursor = m.end()
        append_text(parent, inside[cursor:])
        append_text(parent, "]")
        return True
    m = re.fullmatch(r"(?i)Figure\s+(\d+)", token)
    if m:
        x = etree.SubElement(parent, "xref", {"ref-type": "fig", "rid": f"figure{m.group(1)}"})
        x.text = token
        return True
    m = re.fullmatch(r"(?i)Table\s+(\d+)", token)
    if m:
        x = etree.SubElement(parent, "xref", {"ref-type": "table", "rid": f"table{m.group(1)}"})
        x.text = token
        return True
    m = re.fullmatch(r"(?i)Multimedia Appendix\s+(\d+)", token)
    if m:
        x = etree.SubElement(parent, "xref", {"ref-type": "supplementary-material", "rid": f"app{m.group(1)}"})
        x.text = token
        return True
    m = re.fullmatch(r"(?i)Checklist\s+(\d+)", token)
    if m:
        # JMIR sequence treats Checklist 1 as the next supplementary item after Multimedia Appendices.
        num = m.group(1)
        x = etree.SubElement(parent, "xref", {"ref-type": "supplementary-material", "rid": f"app{num}"})
        x.text = token
        x.set("data-checklist", "yes")  # temporary marker fixed after app inventory known
        return True
    return False


def append_text_with_xrefs(parent: etree._Element, text: str) -> None:
    pos = 0
    for m in XREF_PATTERN.finditer(text):
        append_text(parent, text[pos:m.start()])
        append_xref_token(parent, m.group())
        pos = m.end()
    append_text(parent, text[pos:])


def append_span(parent: etree._Element, span: Span) -> None:
    container = parent
    if span.href:
        link = etree.SubElement(container, "ext-link", {"ext-link-type": "uri"})
        set_xlink(link, "href", span.href)
        container = link
    if span.bold:
        container = etree.SubElement(container, "bold")
    if span.italic:
        container = etree.SubElement(container, "italic")
    if span.sup:
        container = etree.SubElement(container, "sup")
    if span.sub:
        container = etree.SubElement(container, "sub")
    append_text_with_xrefs(container, span.text)


def append_paragraph_inline(parent: etree._Element, p: Paragraph, strip_prefix: Optional[str] = None) -> None:
    spans = paragraph_spans(p)
    if strip_prefix:
        remaining = len(strip_prefix)
        new_spans = []
        for s in spans:
            if remaining <= 0:
                new_spans.append(s); continue
            if len(s.text) <= remaining:
                remaining -= len(s.text); continue
            s = Span(s.text[remaining:], s.bold, s.italic, s.sup, s.sub, s.href)
            remaining = 0
            new_spans.append(s)
        spans = new_spans
    for s in spans:
        append_span(parent, s)


# -------------------------- reference parser --------------------------

ID_LABELS = (("doi", "doi"), ("medline", "Medline"), ("pmid", "PMID"), ("pmc", "PMCID"))
DOI_URL = re.compile(r"\bhttps?://(?:dx\.)?doi\.org/(\S+)", flags=re.I)
URL = re.compile(r"\bhttps?://[^\s\]]+|\bwww\.[^\s\]]+", flags=re.I)
ACCESSED = re.compile(
    r"\[\s*accessed\s+([^\]]+?)\s*\]"
    r"|\baccessed\s+(?:on\s+)?(\d{4}-\d{2}-\d{2}|[A-Za-z]+\.?\s+\d{1,2},?\s+\d{4}|\d{1,2}\s+[A-Za-z]+\s+\d{4})\.?",
    flags=re.I,
)
# "Surname AB" tokens: "Epstein RM", "O’Keeffe M", "Ferreira ML", optionally ending in "et al".
_AUTHOR = r"[A-Z][^\s,.;:]*(?:[ -][A-Z][^\s,.;:]*)* [A-Z]{1,4}"
AUTHOR_LIST = re.compile(rf"^{_AUTHOR}(?:,\s*(?:{_AUTHOR}|et al))*$")


def _cut(text: str, m: re.Match) -> str:
    return text[:m.start()] + " " + text[m.end():]


def extract_ids(text: str) -> tuple[str, dict[str, str]]:
    """Remove identifiers in bracketed ([doi: X]) or bare (doi: X) form and return them by pub-id-type."""
    ids = {}
    for typ, label in ID_LABELS:
        m = re.search(rf"\[\s*{label}:\s*([^\]]+)\]|\b{label}:\s*(\S+)", text, flags=re.I)
        if m:
            ids[typ] = (m.group(1) or m.group(2)).strip().rstrip(".,;")
            text = _cut(text, m)
    if "doi" not in ids:
        m = DOI_URL.search(text)
        if m:
            ids["doi"] = m.group(1).rstrip(".,;")
            text = _cut(text, m)
    m = re.search(r"\[?\bISBN:?\s*([0-9Xx][0-9Xx-]+)\]?", text)
    if m:
        ids["other"] = m.group(1).replace("-", "")
        text = _cut(text, m)
    return normalize_space(text), ids


def extract_web(text: str, links: Sequence[tuple[str, str]]) -> tuple[str, Optional[str], Optional[str]]:
    """Remove a trailing access date and URL; `links` are the paragraph's (display text, href) hyperlinks."""
    accessed = None
    m = ACCESSED.search(text)
    if m:
        accessed = (m.group(1) or m.group(2)).strip()
        text = _cut(text, m)
    url = None
    m = URL.search(text)
    if m:
        url = m.group().rstrip(".,;")
        text = _cut(text, m)
    text = normalize_space(text)
    for shown, href in links:
        shown = shown.strip()
        if not shown or shown not in text or "doi.org/" in href:
            continue
        url = url or href
        # Link labels such as "SciELO Link" at the end of a reference stand in for the URL.
        if text.rstrip(" .").endswith(shown):
            text = normalize_space(text.rstrip(" .")[: -len(shown)])
    return text, url, accessed


def _add_pages(cit: etree._Element, pages: str) -> None:
    pages = pages.strip()
    if re.fullmatch(r"[^-–]+[-–][^-–]+", pages):
        fp, lp = re.split(r"[-–]", pages, maxsplit=1)
        etree.SubElement(cit, "fpage").text = fp.strip()
        etree.SubElement(cit, "lpage").text = lp.strip()
    else:
        etree.SubElement(cit, "fpage").text = pages


def _add_publisher(cit: etree._Element, publisher: str) -> None:
    """'Philadelphia, PA: Elsevier' -> publisher-loc + publisher-name."""
    loc, sep, name = publisher.rpartition(": ")
    if sep and loc.strip():
        etree.SubElement(cit, "publisher-loc").text = loc.strip()
    etree.SubElement(cit, "publisher-name").text = (name if sep else publisher).strip()


def parse_journal_reference(text: str, cit: etree._Element) -> bool:
    if ". " not in text:
        return False
    authors, rest = text.split(". ", 1)

    # Find bibliographic tail: optional Month Day, then year;volume(issue):pages.
    pat = re.compile(
        r"(?:(?P<mon>[A-Za-z]{3,9})\s+(?:(?P<day>\d{1,2}),\s*)?)?"
        r"(?P<year>\d{4});(?P<vol>[^(:;\.]+)"
        r"(?:\((?P<issue>[^)]+)\))?:(?P<pages>[^\.]+)\.?$"
    )
    m = pat.search(rest)
    if not m:
        return False
    pre = rest[:m.start()].rstrip(". ")
    if ". " not in pre:
        return False
    title, source = pre.rsplit(". ", 1)
    add_person_group(cit, authors, "author")
    etree.SubElement(cit, "article-title").text = title.strip()
    etree.SubElement(cit, "source").text = source.strip()
    etree.SubElement(cit, "year").text = m.group("year")
    if m.group("mon"):
        etree.SubElement(cit, "month").text = MONTHS.get(m.group("mon").lower(), m.group("mon"))
    if m.group("day"):
        etree.SubElement(cit, "day").text = str(int(m.group("day")))
    etree.SubElement(cit, "volume").text = m.group("vol").strip()
    if m.group("issue"):
        etree.SubElement(cit, "issue").text = m.group("issue").strip()
    _add_pages(cit, m.group("pages"))
    return True


def parse_book_reference(text: str, cit: etree._Element) -> bool:
    if ". " not in text:
        return False
    authors, rest = text.split(". ", 1)

    # Chapter in a book: Chapter. In: [Editors, editors.] Book. [Loc: ]Publisher; Year[:pages].
    if ". In: " in rest:
        chapter, tail = rest.split(". In: ", 1)
        m = re.match(r"(.+?),\s*editors?\.\s*(.+)$", tail, flags=re.I)
        editors, pubpart = m.groups() if m else (None, tail)
        mm = re.match(r"(.+)\.\s*([^.;]+);\s*(\d{4})(?::([^\.]+))?\.?$", pubpart)
        if not mm:
            return False
        source, publisher, year, pages = mm.groups()
        add_person_group(cit, authors, "author")
        if editors:
            add_person_group(cit, editors, "editor")
        etree.SubElement(cit, "article-title").text = chapter.strip()
        etree.SubElement(cit, "source").text = source.strip()
        etree.SubElement(cit, "year").text = year
        _add_publisher(cit, publisher)
        if pages:
            _add_pages(cit, pages)
    else:
        # Whole book: Title. [1st ed.] [Loc: ]Publisher; Year.
        mm = re.match(r"(.+?)\.\s*(?:(\d+)(?:st|nd|rd|th)\s+ed\.\s*)?([^.;]+);\s*(\d{4})\.?$", rest, flags=re.I)
        if not mm:
            return False
        source, edition, publisher, year = mm.groups()
        add_person_group(cit, authors, "author")
        etree.SubElement(cit, "source").text = source.strip()
        etree.SubElement(cit, "year").text = year
        if edition:
            etree.SubElement(cit, "edition").text = edition
        _add_publisher(cit, publisher)
    return True


def parse_web_reference(text: str, cit: etree._Element) -> bool:
    """[Authors.] [Title.] Site[. [Publisher;] Year]. Only called when a URL was found."""
    parts = [s.strip() for s in re.split(r"\.\s+", text.rstrip(". ")) if s.strip()]
    if not parts:
        return False
    authors = parts.pop(0) if len(parts) > 1 and AUTHOR_LIST.match(parts[0]) else None
    year = publisher = None
    if len(parts) > 1:
        m = re.fullmatch(r"(?:(.+?)[;,]\s*)?(\d{4})", parts[-1])
        if m:
            publisher, year = m.groups()
            parts.pop()
    if authors:
        add_person_group(cit, authors, "author")
    if len(parts) > 1:
        etree.SubElement(cit, "article-title").text = ". ".join(parts[:-1])
    etree.SubElement(cit, "source").text = parts[-1]
    if year:
        etree.SubElement(cit, "year").text = year
    if publisher:
        _add_publisher(cit, publisher)
    return True


def _paragraph_links(p: Paragraph) -> list[tuple[str, str]]:
    return [(s.text, s.href) for s in paragraph_spans(p) if s.href]


def build_reference(parent: etree._Element, p: Paragraph, warnings: list[str]) -> None:
    raw = normalize_space(p.text.replace("\t", " "))
    m = re.match(r"(\d+)\.\s*(.*)", raw)
    num = m.group(1) if m else str(len(parent.findall("ref")) + 1)
    body = body_raw = m.group(2) if m else raw
    ref = etree.SubElement(parent, "ref", {"id": f"ref{num}"})
    etree.SubElement(ref, "label").text = num

    # Identifiers, URLs and access dates can appear in any order after the
    # bibliographic core; strip them first so the core parsers see only the citation.
    body, ids = extract_ids(body)
    body, url, accessed = extract_web(body, _paragraph_links(p))

    # Classification: edited chapter or whole book markers, otherwise journal.
    is_book = bool(re.search(r"\bIn:\s|\b\d+(?:st|nd|rd|th)\s+ed\.\s", body)) or "other" in ids
    # Also whole books typically have Publisher; YEAR and no ;volume:pages pattern.
    if not is_book and re.search(r"\.\s*[^.;]+;\s*\d{4}\.?$", body) and not re.search(r"\d{4};[^:]+:", body):
        is_book = True
    parsers = [("book", parse_book_reference), ("journal", parse_journal_reference)]
    if not is_book:
        parsers.reverse()
    if url:
        parsers.append(("web", parse_web_reference))

    for ctype, parse in parsers:
        cit = etree.Element("nlm-citation", {"citation-type": ctype})
        if parse(body, cit):
            ref.append(cit)
            if url:
                comment = etree.SubElement(cit, "comment")
                link = etree.SubElement(comment, "ext-link", {"ext-link-type": "uri"})
                set_xlink(link, "href", url)
                link.text = url
                if accessed:
                    link.tail = f" [Accessed {accessed}]"
            for typ, val in ids.items():
                etree.SubElement(cit, "pub-id", {"pub-id-type": typ}).text = val
            return

    # Structural fallback is lossless but flagged for human review. NLM 2.0 has no
    # mixed-citation (added in 3.0); its unstructured equivalent is <citation>.
    fallback = etree.SubElement(ref, "citation")
    fallback.text = body_raw
    warnings.append(f"Reference {num}: structured parse failed; emitted unstructured citation")


# ------------------------------ converter ------------------------------

@dataclass
class ConversionContext:
    profile: dict
    warnings: list[str] = field(default_factory=list)
    volume: str = ""
    issue: str = "1"
    elocation: str = ""
    doi: str = ""
    self_url: str = ""
    pub_date: Optional[tuple[str, str, str]] = None
    submitted: Optional[tuple[str, str, str]] = None
    revised: Optional[tuple[str, str, str]] = None
    accepted: Optional[tuple[str, str, str]] = None
    editor: Optional[str] = None
    reviewers: list[str] = field(default_factory=list)
    app_count: int = 0


class ManuscriptToJATS:
    def __init__(self, docx: Path, profile: dict):
        self.docx_path = docx
        self.doc = Document(str(docx))
        self.profile = profile
        self.ctx = ConversionContext(profile=profile)
        self.blocks = list(iter_block_items(self.doc))
        self.paragraphs = [x for x in self.blocks if isinstance(x, Paragraph)]
        self.styles = profile["styles"]
        self._extract_production_metadata()
        self.root = etree.Element(
            "article",
            nsmap={"mml": MML, "xlink": XLINK},
            attrib={
                "dtd-version": "2.0",
                f"{{{XML_NS}}}lang": profile.get("language", "en"),
                "article-type": profile.get("article_type", "research-article"),
            },
        )
        self.front = etree.SubElement(self.root, "front")
        self.body = etree.SubElement(self.root, "body")
        self.back = etree.SubElement(self.root, "back")

    def _style_is(self, p: Paragraph, role: str) -> bool:
        return style_name(p) in self.styles.get(role, [])

    def _find_first(self, role: str) -> Optional[Paragraph]:
        return next((p for p in self.paragraphs if self._style_is(p, role)), None)

    def _extract_production_metadata(self) -> None:
        alltext = "\n".join(p.text for p in self.paragraphs)
        # Citation block: JMIR Med Educ 2026;12:e95904
        abbr = re.escape(self.profile["journal"]["abbrev_title"])
        m = re.search(rf"{abbr}\s+(\d{{4}});(\d+):e?(\d+)", alltext)
        if m:
            _, self.ctx.volume, eid = m.groups()
            self.ctx.elocation = eid
        m = re.search(r"(?im)^doi:\s*([^\s]+)", alltext)
        if m:
            self.ctx.doi = m.group(1).strip()
        m = re.search(r"(?im)^URL:\s*(https?://\S+)", alltext)
        if m:
            self.ctx.self_url = m.group(1).strip()
        self.ctx.issue = self.profile["journal"].get("issue", "1")

        # Editorial line may span multiple paragraphs.
        edtext = " ".join(p.text.strip() for p in self.paragraphs[-20:])
        m = re.search(r"Edited by\s+([^;]+)", edtext, re.I)
        if m:
            self.ctx.editor = normalize_space(m.group(1))
        m = re.search(r"peer-reviewed by\s+([^;]+)", edtext, re.I)
        if m:
            self.ctx.reviewers = [normalize_space(x) for x in m.group(1).split(",")]
        for label, attr in [
            ("submitted", "submitted"),
            ("final revised version received", "revised"),
            ("revised version received", "revised"),
            ("accepted", "accepted"),
            ("published", "pub_date"),
        ]:
            m = re.search(rf"{re.escape(label)}\s+([^;]+)", edtext, re.I)
            if m and not getattr(self.ctx, attr):
                parts = date_parts(m.group(1))
                if parts:
                    setattr(self.ctx, attr, parts)

    # -------- front --------
    def build_front(self) -> None:
        j = self.profile["journal"]
        jm = etree.SubElement(self.front, "journal-meta")
        etree.SubElement(jm, "journal-id", {"journal-id-type": "nlm-ta"}).text = j["nlm_ta"]
        etree.SubElement(jm, "journal-id", {"journal-id-type": "publisher-id"}).text = j["publisher_id"]
        if j.get("index"):
            etree.SubElement(jm, "journal-id", {"journal-id-type": "index"}).text = str(j["index"])
        etree.SubElement(jm, "journal-title").text = j["title"]
        etree.SubElement(jm, "abbrev-journal-title").text = j["abbrev_title"]
        etree.SubElement(jm, "issn", {"pub-type": "epub"}).text = j["issn_epub"]
        pub = etree.SubElement(jm, "publisher")
        etree.SubElement(pub, "publisher-name").text = j["publisher_name"]
        etree.SubElement(pub, "publisher-loc").text = j["publisher_location"]

        am = etree.SubElement(self.front, "article-meta")
        if self.ctx.elocation:
            pid = f"v{self.ctx.volume}i{self.ctx.issue}e{self.ctx.elocation}"
            etree.SubElement(am, "article-id", {"pub-id-type": "publisher-id"}).text = pid
        if self.ctx.doi:
            etree.SubElement(am, "article-id", {"pub-id-type": "doi"}).text = self.ctx.doi

        catp = self._find_first("article_type")
        cats = etree.SubElement(am, "article-categories")
        sg = etree.SubElement(cats, "subj-group", {"subj-group-type": "heading"})
        etree.SubElement(sg, "subject").text = catp.text.strip() if catp else "Original Paper"

        titlep = self._find_first("article_title")
        tg = etree.SubElement(am, "title-group")
        title = etree.SubElement(tg, "article-title")
        if titlep:
            append_paragraph_inline(title, titlep)
        else:
            self.ctx.warnings.append("Article title not found")

        self._build_authors(am)
        self._build_affiliations(am)
        self._build_editor_reviewers(am)
        self._build_author_notes(am)
        self._build_dates(am)
        self._build_copyright_license(am)
        self._build_abstract_keywords(am)

    def _build_authors(self, am: etree._Element) -> None:
        p = self._find_first("authors")
        if not p:
            self.ctx.warnings.append("Author paragraph not found")
            return
        # Use superscript runs as deterministic affiliation markers.
        token = ""
        for sp in paragraph_spans(p):
            if sp.sup:
                token += f"<SUP>{sp.text}</SUP>"
            else:
                token += sp.text
        chunks = [x.strip() for x in token.split(";") if x.strip()]
        cg = etree.SubElement(am, "contrib-group")
        corresp_lines = [x.text.strip() for x in self.paragraphs if self._style_is(x, "correspondence")]
        corresp_name = corresp_lines[1].split(",",1)[0].strip() if len(corresp_lines) > 1 else ""
        for chunk in chunks:
            m = re.match(r"(.+?)(?:<SUP>(.*?)</SUP>)?\s*,\s*(.+)$", chunk)
            if not m:
                self.ctx.warnings.append(f"Could not parse author: {chunk}")
                continue
            name_txt, sup, degrees = m.groups()
            attrs = {"contrib-type": "author"}
            if corresp_name and normalize_space(name_txt) == normalize_space(corresp_name):
                attrs["corresp"] = "yes"
            if sup and "*" in sup:
                attrs["equal-contrib"] = "yes"
            c = etree.SubElement(cg, "contrib", attrs)
            person_name(c, name_txt)
            etree.SubElement(c, "degrees").text = normalize_space(degrees)
            if sup:
                for n in re.findall(r"\d+", sup):
                    x = etree.SubElement(c, "xref", {"ref-type": "aff", "rid": f"aff{n}"})
                    x.text = n
                if "*" in sup:
                    x = etree.SubElement(c, "xref", {"ref-type": "fn", "rid": "equal-contrib1"})
                    x.text = "*"

    def _split_affiliation(self, text: str) -> tuple[str, list[str], str]:
        parts = [x.strip() for x in text.split(",") if x.strip()]
        country = ""
        if parts and parts[-1] in self.profile.get("country_names", []):
            country = parts.pop()
        # Heuristic: institution ends at last segment containing a strong organization keyword.
        org_keywords = ("University", "School", "College", "Institute", "Hospital", "Center", "Centre", "System")
        end = -1
        for i, part in enumerate(parts):
            if any(k.lower() in part.lower() for k in org_keywords):
                end = i
        if end >= 0:
            institution = ", ".join(parts[:end+1])
            addr = parts[end+1:]
        else:
            institution = parts[0] if parts else text
            addr = parts[1:]
        return institution, addr, country

    def _build_affiliations(self, am: etree._Element) -> None:
        for p in self.paragraphs:
            if not self._style_is(p, "affiliation"):
                continue
            t = p.text.strip()
            if t.startswith("*"):
                continue
            m = re.match(r"(\d+)\s*(.*)", t)
            if not m:
                continue
            num, rest = m.groups()
            aff = etree.SubElement(am, "aff", {"id": f"aff{num}"})
            institution, addr, country = self._split_affiliation(rest)
            # JMIR commonly repeats the corresponding author's street address in
            # the linked affiliation even when the Word affiliation line omits it.
            # If the corresponding author uses this affiliation, recover a clear
            # street-address line from the correspondence block without inventing data.
            authorp = self._find_first("authors")
            corr_lines = [q.text.strip() for q in self.paragraphs if self._style_is(q, "correspondence")]
            corr_name = corr_lines[1].split(",", 1)[0].strip() if len(corr_lines) > 1 else ""
            corr_affs = set()
            if authorp and corr_name:
                tok = ""
                for sp in paragraph_spans(authorp):
                    tok += f"<SUP>{sp.text}</SUP>" if sp.sup else sp.text
                for chunk in tok.split(";"):
                    if corr_name in chunk:
                        sm = re.search(r"<SUP>(.*?)</SUP>", chunk)
                        if sm:
                            corr_affs.update(re.findall(r"\d+", sm.group(1)))
            if num in corr_affs:
                street = next((x for x in corr_lines if re.search(r"(?i)\b(street|st\.?|road|rd\.?|avenue|ave\.?|boulevard|blvd\.?|lane|ln\.?)\b", x)), None)
                if street and street not in addr and street not in institution:
                    addr.insert(0, street)
            etree.SubElement(aff, "institution").text = institution
            for a in addr:
                etree.SubElement(aff, "addr-line").text = a
            if country:
                etree.SubElement(aff, "country").text = country

    def _build_editor_reviewers(self, am: etree._Element) -> None:
        if self.ctx.editor:
            cg = etree.SubElement(am, "contrib-group")
            c = etree.SubElement(cg, "contrib", {"contrib-type": "editor"})
            person_name(c, self.ctx.editor)
        if self.ctx.reviewers:
            cg = etree.SubElement(am, "contrib-group")
            for r in self.ctx.reviewers:
                c = etree.SubElement(cg, "contrib", {"contrib-type": "reviewer"})
                person_name(c, r)

    def _build_author_notes(self, am: etree._Element) -> None:
        lines = [p.text.strip() for p in self.paragraphs if self._style_is(p, "correspondence")]
        equal = next((p.text.strip() for p in self.paragraphs if self._style_is(p, "affiliation") and p.text.strip().startswith("*")), None)
        if not lines and not equal:
            return
        an = etree.SubElement(am, "author-notes")
        if lines:
            corr = etree.SubElement(an, "corresp")
            text = ", ".join([x for x in lines[1:] if x and not x.lower().startswith("email:")])
            email = next((x.split(":",1)[1].strip() for x in lines if x.lower().startswith("email:")), "")
            corr.text = "Correspondence to " + text
            if email:
                corr.text = (corr.text or "") + "; "
                etree.SubElement(corr, "email").text = email
        if equal:
            fn = etree.SubElement(an, "fn", {"fn-type": "equal", "id": "equal-contrib1"})
            etree.SubElement(fn, "label").text = "*"
            etree.SubElement(fn, "p").text = equal.lstrip("*").strip()

    def _build_dates(self, am: etree._Element) -> None:
        if self.ctx.pub_date:
            _, _, y = self.ctx.pub_date
            pd = etree.SubElement(am, "pub-date", {"pub-type": "collection"})
            etree.SubElement(pd, "year").text = y
            pd = etree.SubElement(am, "pub-date", {"pub-type": "epub"})
            d,m,y = self.ctx.pub_date
            etree.SubElement(pd, "day").text = str(int(d))
            etree.SubElement(pd, "month").text = str(int(m))
            etree.SubElement(pd, "year").text = y
        if self.ctx.volume:
            etree.SubElement(am, "volume").text = self.ctx.volume
        if self.ctx.elocation:
            etree.SubElement(am, "elocation-id").text = "e" + self.ctx.elocation
        hist = etree.SubElement(am, "history")
        for typ, parts in (("received", self.ctx.submitted), ("rev-recd", self.ctx.revised), ("accepted", self.ctx.accepted)):
            if parts:
                add_date(hist, typ, parts)

    def _build_copyright_license(self, am: etree._Element) -> None:
        cp = next((p for p in self.paragraphs if p.text.strip().startswith("©")), None)
        if not cp:
            return
        text = normalize_space(cp.text)
        # split before license sentence
        marker = "This is an open-access article"
        before, sep, after = text.partition(marker)
        cs = etree.SubElement(am, "copyright-statement")
        ctext = before.strip()
        pos = 0
        for m in re.finditer(r"https?://[^\s,)]+/?", ctext):
            append_text(cs, ctext[pos:m.start()])
            url = m.group(0)
            link = etree.SubElement(cs, "ext-link", {"ext-link-type": "uri"})
            set_xlink(link, "href", url)
            link.text = url
            pos = m.end()
        append_text(cs, ctext[pos:])
        year = self.ctx.pub_date[2] if self.ctx.pub_date else re.search(r"20\d{2}", text).group(0)
        etree.SubElement(am, "copyright-year").text = year
        lic = etree.SubElement(am, "license", {"license-type": "open-access"})
        set_xlink(lic, "href", "https://creativecommons.org/licenses/by/4.0/")
        lp = etree.SubElement(lic, "p")
        lictext = (marker + after).strip() if sep else text
        # URLs become ext-link nodes.
        pos = 0
        for m in re.finditer(r"https?://[^\s,)]+/?", lictext):
            append_text(lp, lictext[pos:m.start()])
            url = m.group(0)
            link = etree.SubElement(lp, "ext-link", {"ext-link-type": "uri"})
            set_xlink(link, "href", url)
            link.text = url
            pos = m.end()
        append_text(lp, lictext[pos:])
        if self.ctx.self_url:
            su = etree.SubElement(am, "self-uri")
            set_xlink(su, "type", "simple")
            set_xlink(su, "href", self.ctx.self_url)

    def _build_abstract_keywords(self, am: etree._Element) -> None:
        absps = [p for p in self.paragraphs if self._style_is(p, "abstract")]
        abstract = etree.SubElement(am, "abstract")
        for p in absps:
            t = p.text.strip()
            if t.lower() == "abstract":
                continue
            m = re.match(r"([^:]{2,30}):\s*(.*)", t, flags=re.S)
            if not m:
                continue
            label, rest = m.groups()
            secattrs = {}
            lower = label.lower()
            typed = set(self.profile.get("abstract_sec_types", ["methods", "results", "conclusions"]))
            if lower in typed:
                secattrs["sec-type"] = lower
            sec = etree.SubElement(abstract, "sec", secattrs)
            etree.SubElement(sec, "title").text = label
            pp = etree.SubElement(sec, "p")
            # Reuse run formatting while dropping "Label: " prefix.
            append_paragraph_inline(pp, p, strip_prefix=t[:t.find(":")+1] + (" " if rest and t[t.find(":")+1:].startswith(" ") else ""))
        kwp = self._find_first("keywords")
        if kwp:
            kg = etree.SubElement(am, "kwd-group")
            txt = re.sub(r"^Keywords:\s*", "", kwp.text.strip(), flags=re.I)
            for k in [x.strip() for x in txt.split(";") if x.strip()]:
                etree.SubElement(kg, "kwd").text = k

    # -------- body/back --------
    def build_body_and_back(self) -> None:
        start = False
        current_top: Optional[etree._Element] = None
        current_sec: Optional[etree._Element] = None
        top_by_title: dict[str, etree._Element] = {}
        sec_counter = 0
        sub_counters: dict[str, int] = {}
        pending_table: Optional[etree._Element] = None
        current_back_heading = ""
        notes_el: Optional[etree._Element] = None
        fn_group: Optional[etree._Element] = None
        glossary: Optional[etree._Element] = None
        def_list: Optional[etree._Element] = None
        ref_list: Optional[etree._Element] = None
        app_group: Optional[etree._Element] = None
        current_app: Optional[etree._Element] = None
        current_app_num: Optional[int] = None
        figure_num = 0
        table_num = 0

        back_headings = {
            self.profile["back"]["acknowledgments"],
            *self.profile["back"]["notes"],
            self.profile["back"]["contributors"],
            self.profile["back"]["conflicts"],
            self.profile["back"]["appendices"],
            self.profile["back"]["references"],
            self.profile["back"]["abbreviations"],
        }
        appendix_item = self.profile["back"].get("appendix_item_pattern", r"(?:Multimedia Appendix|Checklist)\s+\d+")

        for block in self.blocks:
            if isinstance(block, Table):
                if not start or current_back_heading:
                    continue
                if pending_table is None:
                    self.ctx.warnings.append("DOCX table encountered without preceding TableCaption")
                    continue
                self._populate_table(pending_table, block)
                continue

            p = block
            st = style_name(p)
            text = p.text.strip()
            if not text:
                continue
            if not start:
                if self._style_is(p, "heading1") and text == "Introduction":
                    start = True
                else:
                    continue

            # A single appendix heading ("Multimedia Appendix 1", "Checklist 1") opens an
            # app-group item whether it sits under a "Multimedia Appendices" Head1 as a
            # Head2, or stands alone as its own Head1 after the other back matter.
            if (
                current_back_heading
                and (self._style_is(p, "heading1") or self._style_is(p, "heading2"))
                and re.fullmatch(appendix_item, text, flags=re.I)
            ):
                current_back_heading = self.profile["back"]["appendices"]
                if app_group is None:
                    app_group = etree.SubElement(self.back, "app-group")
                self.ctx.app_count += 1
                current_app_num = self.ctx.app_count
                current_app = etree.SubElement(app_group, "supplementary-material", {"id": f"app{current_app_num}"})
                etree.SubElement(current_app, "label").text = text
                continue

            # Heading boundary into back matter.
            if (self._style_is(p, "heading1") or self._style_is(p, "references_heading")) and text in back_headings:
                current_back_heading = text
                current_top = current_sec = None
                pending_table = None
                if text == self.profile["back"]["acknowledgments"]:
                    if self.back.find("ack") is None:
                        etree.SubElement(self.back, "ack")
                elif text in self.profile["back"]["notes"]:
                    if notes_el is None:
                        notes_el = etree.SubElement(self.back, "notes")
                    s = etree.SubElement(notes_el, "sec")
                    etree.SubElement(s, "title").text = text
                elif text in {self.profile["back"]["contributors"], self.profile["back"]["conflicts"]}:
                    if fn_group is None:
                        fn_group = etree.SubElement(self.back, "fn-group")
                    typ = "con" if text == self.profile["back"]["contributors"] else "conflict"
                    etree.SubElement(fn_group, "fn", {"fn-type": typ})
                elif text == self.profile["back"]["abbreviations"]:
                    glossary = etree.SubElement(self.back, "glossary")
                    etree.SubElement(glossary, "title").text = text
                    def_list = etree.SubElement(glossary, "def-list")
                elif text == self.profile["back"]["references"]:
                    ref_list = etree.SubElement(self.back, "ref-list")
                    etree.SubElement(ref_list, "title").text = text
                elif text == self.profile["back"]["appendices"]:
                    if app_group is None:
                        app_group = etree.SubElement(self.back, "app-group")
                continue

            # Body sections.
            if not current_back_heading:
                if self._style_is(p, "heading1"):
                    parent_override = self.profile["body"].get("parent_overrides", {}).get(text)
                    if parent_override and parent_override in top_by_title:
                        parent = top_by_title[parent_override]
                        topid = parent.get("id")
                        sub_counters[topid] = sub_counters.get(topid, 0) + 1
                        current_sec = etree.SubElement(parent, "sec", {"id": f"{topid}-{sub_counters[topid]}"})
                        etree.SubElement(current_sec, "title").text = text
                        current_top = parent
                    else:
                        sec_counter += 1
                        attrs = {"id": f"s{sec_counter}"}
                        if text in self.profile["body"].get("sec_type", {}):
                            attrs["sec-type"] = self.profile["body"]["sec_type"][text]
                        current_top = etree.SubElement(self.body, "sec", attrs)
                        current_sec = current_top
                        top_by_title[text] = current_top
                        etree.SubElement(current_top, "title").text = text
                    continue
                if self._style_is(p, "heading2"):
                    if current_top is None:
                        self.ctx.warnings.append(f"Heading2 outside body section: {text}")
                        continue
                    topid = current_top.get("id")
                    sub_counters[topid] = sub_counters.get(topid, 0) + 1
                    current_sec = etree.SubElement(current_top, "sec", {"id": f"{topid}-{sub_counters[topid]}"})
                    etree.SubElement(current_sec, "title").text = text
                    continue
                if self._style_is(p, "body"):
                    if current_sec is None:
                        continue
                    sec_title = current_sec.findtext("title") or ""
                    merge_sections = set(self.profile.get("merge_paragraphs_sections", []))
                    existing = current_sec.findall("./p")
                    if sec_title in merge_sections and existing:
                        pp = existing[-1]
                        append_text(pp, " ")
                    else:
                        pp = etree.SubElement(current_sec, "p")
                    append_paragraph_inline(pp, p)
                    continue
                if self._style_is(p, "figure_caption"):
                    if current_sec is None:
                        continue
                    figure_num += 1
                    self._add_figure(current_sec, p, figure_num)
                    continue
                if self._style_is(p, "table_caption"):
                    if current_sec is None:
                        continue
                    table_num += 1
                    pending_table = self._add_table_wrap(current_sec, p, table_num)
                    continue
                if self._style_is(p, "table_note") and pending_table is not None:
                    self._add_table_note(pending_table, p, table_num)
                    pending_table = None
                    continue
                continue

            # Back matter content.
            if current_back_heading == self.profile["back"]["acknowledgments"]:
                if self._style_is(p, "body"):
                    ack = self.back.find("ack")
                    pp = etree.SubElement(ack, "p")
                    append_paragraph_inline(pp, p)
                continue

            if current_back_heading in self.profile["back"]["notes"]:
                if self._style_is(p, "body"):
                    sec = notes_el.findall("sec")[-1]
                    pp = etree.SubElement(sec, "p")
                    append_paragraph_inline(pp, p)
                continue

            if current_back_heading in {self.profile["back"]["contributors"], self.profile["back"]["conflicts"]}:
                if self._style_is(p, "list") or self._style_is(p, "body"):
                    fn = fn_group.findall("fn")[-1]
                    pp = etree.SubElement(fn, "p")
                    append_paragraph_inline(pp, p)
                continue

            if current_back_heading == self.profile["back"]["appendices"]:
                if current_app is not None and (self._style_is(p, "list") or self._style_is(p, "body")):
                    # "[PDF File (Adobe File), 132 KB-Checklist 1]" is the file line; anything
                    # else before it is the description.
                    if SUPP_FILE_LINE.match(text) or current_app.find("p") is not None:
                        if current_app.find("media") is None:
                            self._add_supp_media(current_app, p, current_app_num)
                    else:
                        pp = etree.SubElement(current_app, "p")
                        append_paragraph_inline(pp, p)
                continue

            if current_back_heading == self.profile["back"]["references"]:
                if ref_list is not None and self._style_is(p, "reference"):
                    build_reference(ref_list, p, self.ctx.warnings)
                continue

            if current_back_heading == self.profile["back"]["abbreviations"]:
                if def_list is not None and self._style_is(p, "list"):
                    m = re.match(r"([^:]+):\s*(.*)", text)
                    if m:
                        term, definition = m.groups()
                        di = etree.SubElement(def_list, "def-item")
                        t = etree.SubElement(di, "term", {"id": f"abb{len(def_list.findall('def-item'))}"})
                        t.text = term
                        d = etree.SubElement(di, "def")
                        etree.SubElement(d, "p").text = definition
                continue

        self._fix_checklist_xrefs()

    def _add_figure(self, parent: etree._Element, p: Paragraph, num: int) -> None:
        text = p.text.strip()
        m = re.match(r"Figure\s+(\d+)\.\s*(.*)", text, flags=re.I|re.S)
        n = int(m.group(1)) if m else num
        caption = m.group(2) if m else text
        fig = etree.SubElement(parent, "fig", {"position": "float", "id": f"figure{n}"})
        etree.SubElement(fig, "label").text = f"Figure {n}."
        cap = etree.SubElement(fig, "caption")
        cp = etree.SubElement(cap, "p")
        append_text_with_xrefs(cp, caption)
        a = self.profile["assets"]
        filename = a["figure_pattern"].format(
            journal=self.profile["journal"]["publisher_id"], volume=self.ctx.volume,
            issue=self.ctx.issue, elocation=self.ctx.elocation, num=n, ext=a.get("figure_extension", ".png")
        )
        g = etree.SubElement(fig, "graphic", {"mimetype": "image", "position": "float"})
        set_xlink(g, "type", "simple")
        set_xlink(g, "href", filename)

    def _add_table_wrap(self, parent: etree._Element, p: Paragraph, num: int) -> etree._Element:
        text = p.text.strip()
        m = re.match(r"Table\s+(\d+)\.\s*(.*)", text, flags=re.I|re.S)
        n = int(m.group(1)) if m else num
        caption = m.group(2) if m else text
        tw = etree.SubElement(parent, "table-wrap", {"id": f"t{n}", "position": "float"})
        etree.SubElement(tw, "label").text = f"Table {n}."
        cap = etree.SubElement(tw, "caption")
        cp = etree.SubElement(cap, "p")
        # trailing superscript note marker e.g. '... (ITEM)a.'
        mm = re.match(r"(.*?)([a-z])\.$", caption)
        if mm:
            append_text_with_xrefs(cp, mm.group(1))
            sup = etree.SubElement(cp, "sup")
            x = etree.SubElement(sup, "xref", {"ref-type": "table-fn", "rid": f"table{n}fn1"})
            x.text = mm.group(2)
            append_text(cp, ".")
        else:
            append_text_with_xrefs(cp, caption)
        return tw

    def _populate_table(self, tw: etree._Element, tbl: Table) -> None:
        tnum = re.search(r"\d+", tw.get("id", "1")).group()
        table = etree.SubElement(tw, "table", {"id": f"table{tnum}", "frame": "hsides", "rules": "groups"})
        for r_idx, row in enumerate(tbl.rows):
            container = etree.SubElement(table, "thead" if r_idx == 0 else "tbody") if r_idx in {0,1} else table.find("tbody")
            # For row 0, create thead. For row 1, create tbody once.
            if r_idx == 0:
                parent = container
            else:
                if table.find("tbody") is None:
                    parent = etree.SubElement(table, "tbody")
                else:
                    parent = table.find("tbody")
            tr = etree.SubElement(parent, "tr")
            cells = row.cells
            i = 0
            while i < len(cells):
                cell = cells[i]
                tc = cell._tc
                span = 1
                while i + span < len(cells) and cells[i+span]._tc is tc:
                    span += 1
                attrs = {"align": "left", "valign": "top"}
                if span > 1:
                    attrs["colspan"] = str(span)
                td = etree.SubElement(tr, "td", attrs)
                # target JMIR uses indentation in first column for data/question rows
                cell_text = normalize_space(" ".join(pp.text for pp in cell.paragraphs))
                merged_header = span == len(cells)
                if i == 0 and r_idx > 1 and not merged_header:
                    nc = etree.SubElement(td, "named-content", {"content-type": "indent"})
                    nc.text = "\u00a0\u00a0\u00a0\u00a0"
                for j, pp in enumerate(cell.paragraphs):
                    if j:
                        etree.SubElement(td, "break")
                    # Table cell character styling in the Word manuscript is often
                    # visual-only; JMIR target XML keeps these cells semantically plain.
                    append_text_with_xrefs(td, pp.text)
                i += span

    def _add_table_note(self, tw: etree._Element, p: Paragraph, num: int) -> None:
        foot = etree.SubElement(tw, "table-wrap-foot")
        fn = etree.SubElement(foot, "fn", {"id": f"table{num}fn1"})
        pp = etree.SubElement(fn, "p")
        txt = p.text.strip()
        m = re.match(r"([a-z])\s*(.*)", txt, flags=re.I|re.S)
        if m:
            sup = etree.SubElement(pp, "sup")
            sup.text = m.group(1)
            append_text_with_xrefs(pp, m.group(2))
        else:
            append_paragraph_inline(pp, p)

    def _add_supp_media(self, app: etree._Element, p: Paragraph, num: int) -> None:
        text = p.text.strip().strip("[]")
        # e.g. DOCX File (Microsoft Word File), 26 KB-Multimedia Appendix 1
        m = re.match(r"([A-Z0-9]+)\s+File.*?,\s*(\d+)\s*KB-", text, flags=re.I)
        ext = (m.group(1) if m else "bin").lower()
        kb = m.group(2) if m else ""
        filename = self.profile["assets"]["supp_pattern"].format(
            journal=self.profile["journal"]["publisher_id"], volume=self.ctx.volume,
            issue=self.ctx.issue, elocation=self.ctx.elocation, num=num, ext=ext
        )
        med = etree.SubElement(app, "media")
        set_xlink(med, "href", filename)
        title = f"{ext.upper()} File"
        if kb:
            title += f", {kb} KB"
        set_xlink(med, "title", title)

    def _fix_checklist_xrefs(self) -> None:
        # Checklist n IDs follow all Multimedia Appendix IDs in this manuscript family.
        # Resolve temporary data-checklist markers using actual labels in app-group.
        label_to_id = {}
        for app in self.root.findall(".//app-group/supplementary-material"):
            label_to_id[normalize_space(app.findtext("label") or "").lower()] = app.get("id")
        for x in self.root.findall(".//xref[@data-checklist='yes']"):
            rid = label_to_id.get(normalize_space(x.text or "").lower())
            if rid:
                x.set("rid", rid)
            x.attrib.pop("data-checklist", None)

    def build(self) -> etree._ElementTree:
        self.build_front()
        self.build_body_and_back()
        return etree.ElementTree(self.root)

    def write(self, out: Path) -> None:
        tree = self.build()
        out.parent.mkdir(parents=True, exist_ok=True)
        doctype = self.profile.get("doctype") or '<!DOCTYPE article PUBLIC "-//NLM//DTD Journal Publishing DTD v2.0 20040830//EN" "journalpublishing.dtd">'
        raw_xml = etree.tostring(tree.getroot(), encoding="utf-8", xml_declaration=True, pretty_print=True, doctype=doctype).decode("utf-8")
        out.write_text(encode_non_ascii_entities(raw_xml), encoding="utf-8")


def encode_non_ascii_entities(text: str) -> str:
    """Convert all non-ASCII characters (ord >= 128) to &#xXXXX; hex entities, matching UTF8.pl."""
    text = re.sub(r'&(?!amp;|lt;|gt;|quot;|apos;|#[0-9]+;|#x[0-9a-fA-F]+;)', '&#x0026;', text)
    return re.sub(r'[^\x00-\x7F]', lambda m: f"&#x{ord(m.group(0)):04X};", text)


def convert_docx_to_jats(docx_path: Union[str, Path], profile_path: Union[str, Path]) -> str:
    """Helper to convert a DOCX file into NLM/JATS XML string using the profile."""
    with open(profile_path, "r", encoding="utf-8") as f:
        profile = json.load(f)
    builder = ManuscriptToJATS(Path(docx_path), profile)
    tree = builder.build()
    doctype = profile.get("doctype") or '<!DOCTYPE article PUBLIC "-//NLM//DTD Journal Publishing DTD v2.0 20040830//EN" "journalpublishing.dtd">'
    raw_xml = etree.tostring(tree.getroot(), encoding="utf-8", xml_declaration=True, pretty_print=True, doctype=doctype).decode("utf-8")
    return encode_non_ascii_entities(raw_xml)


def validate_dtd(xml_path: Path, dtd_path: Path) -> tuple[bool, list[str]]:
    parser = etree.XMLParser(load_dtd=False, no_network=True)
    doc = etree.parse(str(xml_path), parser)
    # Load by path, not file object: the DTD pulls in its .ent modules by
    # relative SYSTEM ids, which lxml can only resolve with a base location.
    dtd = etree.DTD(str(dtd_path))
    ok = dtd.validate(doc)
    return ok, [str(e) for e in dtd.error_log]


def main() -> int:
    ap = argparse.ArgumentParser(description="Convert styled JMIR manuscript DOCX to NLM/JATS 2.0 XML")
    ap.add_argument("docx", type=Path)
    ap.add_argument("-p", "--profile", type=Path, required=True)
    ap.add_argument("-o", "--output", type=Path, required=True)
    ap.add_argument("--dtd", type=Path, help="Optional local journalpublishing.dtd for validation")
    ap.add_argument("--warnings", type=Path, help="Optional warning report path")
    args = ap.parse_args()

    profile = json.loads(args.profile.read_text(encoding="utf-8"))
    conv = ManuscriptToJATS(args.docx, profile)
    conv.write(args.output)

    print(f"Wrote: {args.output}")
    print(f"Warnings: {len(conv.ctx.warnings)}")
    for w in conv.ctx.warnings:
        print(" -", w)
    if args.warnings:
        args.warnings.write_text("\n".join(conv.ctx.warnings) + ("\n" if conv.ctx.warnings else ""), encoding="utf-8")

    # Call the perl script and pass the XML file as an argument
        try:
            result = subprocess.run(
                ["perl", perl_script_path, args.output],
                capture_output=True,
                text=True,
                check=True
            )

            print("Perl script executed successfully!")
            print("Output:", result.stdout)

        except subprocess.CalledProcessError as e:
            print("The Perl script failed!")
            print("Error output:", e.stderr)    

    if args.dtd:
        ok, errors = validate_dtd(args.output, args.dtd)
        print("DTD validation:", "PASS" if ok else "FAIL")
        for e in errors:
            print(" -", e)
        if not ok:
            return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
