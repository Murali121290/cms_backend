"""
API router for Ninja Inkflow Language Editing.
Provides endpoints for rule management (CE support JSON sync), job execution, finding reviews, and redlined DOCX export.
"""
import datetime
import logging
import os
import uuid
from typing import Any, Optional
from fastapi import APIRouter, Depends, HTTPException, Body
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app import models
from app.database import get_db
from app.domains.auth.security import get_current_user_from_cookie
from app.domains.processing.language_edit_models import LanguageEditJob, LanguageEditFinding
from app.domains.processing.service import check_permission
from app.models import User

from app.processing.language_editing import docx_io, engine, segmenter
from app.processing.language_editing.rules import load_rules_from_dict, load_house_style_from_dict
from app.services.language_rule_service import (
    get_available_style_profiles,
    get_project_language_rules,
    get_project_language_rules_history,
    save_project_language_rules,
)

# Roles permitted to edit (not just view) the per-project rule selection.
RULES_EDIT_ROLES = {"admin", "projectmanager", "project manager"}

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v2", tags=["Language Editing"])


def _require_language_user(user: Optional[User]) -> User:
    """Require an authenticated user with Language Editing permission.

    Raises 401 if the request has no valid session cookie, 403 if the user's
    roles don't include any member of PROCESS_PERMISSIONS["language"].
    """
    if user is None:
        raise HTTPException(status_code=401, detail="Authentication required.")
    check_permission(user, "language", logger=logger)
    return user


def _require_rules_editor(user: Optional[User]) -> User:
    """Require Admin or Project Manager to edit the per-project rule selection.

    Language Editors can read rules (via `_require_language_user`) but cannot
    change which rules are active for a project.
    """
    if user is None:
        raise HTTPException(status_code=401, detail="Authentication required.")
    user_role_names = {(role.name or "").strip().lower() for role in user.roles}
    if not (user_role_names & RULES_EDIT_ROLES):
        raise HTTPException(
            status_code=403,
            detail=(
                "Only Admin or Project Manager can edit the Rules selection. "
                f"Your roles: {', '.join(sorted(user_role_names)) or 'none'}"
            ),
        )
    return user


# ── Pydantic Request Models ──────────────────────────────────────────────────

class SaveRulesRequest(BaseModel):
    profile_key: Optional[str] = None            # single-style (backward compat)
    profile_keys: Optional[list[str]] = None     # multi-style (new)
    rules: Optional[list[dict[str, Any]]] = None
    variant_to_canonical: Optional[dict[str, str]] = None
    profile_name: Optional[str] = None
    note: Optional[str] = None


class DecisionRequest(BaseModel):
    status: str  # accepted | edited | rejected
    edited_text: Optional[str] = None
    reviewer: Optional[str] = "editor"


# ── Endpoints ────────────────────────────────────────────────────────────────

@router.get("/projects/{project_id}/language-rules")
def get_project_rules(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Fetch current project language rules from CE support along with available base profiles."""
    _require_language_user(current_user)
    active_rules = get_project_language_rules(db, project_id=project_id)
    available_profiles = get_available_style_profiles()
    return {
        "active_rules": active_rules,
        "available_profiles": available_profiles,
    }


@router.post("/projects/{project_id}/language-rules")
def update_project_rules(
    project_id: int,
    body: SaveRulesRequest,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Save/update project language rules, syncing JSON to project's CE support folder.

    Only Admin or Project Manager may change the selection. An audit row is
    written to `project_language_rules_history` capturing the full previous
    and new rule sets plus a summary of which rule IDs were enabled/disabled.
    """
    _require_rules_editor(current_user)
    # Reconcile single vs multi-style payloads (frontend may send either).
    pkeys = body.profile_keys
    if not pkeys:
        pkeys = [body.profile_key] if body.profile_key else ["uk"]
    try:
        updated = save_project_language_rules(
            db,
            project_id=project_id,
            profile_keys=pkeys,
            rules=body.rules,
            variant_to_canonical=body.variant_to_canonical,
            profile_name=body.profile_name,
            changed_by_id=current_user.id,
            note=body.note,
        )
        db.commit()
        return {"ok": True, "rules": updated}
    except ValueError as err:
        db.rollback()
        raise HTTPException(status_code=404, detail=str(err))
    except Exception as exc:
        db.rollback()
        logger.error("Failed to update language rules: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail="Internal server error")


@router.get("/projects/{project_id}/language-rules/history")
def get_project_rules_history(
    project_id: int,
    limit: int = 50,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Return newest-first audit entries for a project's rule changes."""
    _require_language_user(current_user)
    return {
        "project_id": project_id,
        "history": get_project_language_rules_history(db, project_id=project_id, limit=limit),
    }


@router.get("/projects/{project_id}/language-rules/export")
def export_project_rules_json(
    project_id: int,
    download: bool = False,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Serialise the project's rule selection to JSON.

    Default is inline (``application/json``) for the View-JSON modal. Pass
    ``?download=true`` to return an attachment the browser saves as
    ``{project_code}_language_rules.json``.
    """
    import json as _json
    from fastapi.responses import Response

    _require_language_user(current_user)
    cfg = get_project_language_rules(db, project_id=project_id)
    code = cfg.get("project_code") or f"project-{project_id}"
    body = _json.dumps(cfg, indent=2, ensure_ascii=False)
    headers = {}
    if download:
        headers["Content-Disposition"] = f'attachment; filename="{code}_language_rules.json"'
    return Response(content=body, media_type="application/json", headers=headers)


@router.get("/projects/{project_id}/language-rules/export.xlsx")
def export_project_rules_xlsx(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Download the project's rule selection as an Excel workbook.

    Mirrors the per-project JSON config (same source — the CE Support JSON on
    disk). The workbook has one sheet "Language Rules" with columns:
    Rule ID, Message/Description, Category, Severity, Type, Enabled,
    Pattern, Replacement. Enabled rules are highlighted; disabled rules are
    greyed out so the editor can scan the active set at a glance.
    """
    import io
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    _require_language_user(current_user)
    cfg = get_project_language_rules(db, project_id=project_id)
    code = cfg.get("project_code") or f"project-{project_id}"
    rules = cfg.get("rules") or []
    enabled_count = sum(1 for r in rules if r.get("enabled", True))

    wb = Workbook()
    ws = wb.active
    ws.title = "Language Rules"

    # ── Metadata header rows ──────────────────────────────────────────────
    thin = Side(border_style="thin", color="DDDDDD")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    brand_fill = PatternFill("solid", fgColor="312E81")  # indigo-900
    header_fill = PatternFill("solid", fgColor="EEF2FF")
    enabled_fill = PatternFill("solid", fgColor="ECFDF5")  # emerald-50
    disabled_fill = PatternFill("solid", fgColor="F8FAFC")  # slate-50
    sev_fills = {
        "error":      PatternFill("solid", fgColor="FEE2E2"),  # rose-100
        "warning":    PatternFill("solid", fgColor="FEF3C7"),  # amber-100
        "suggestion": PatternFill("solid", fgColor="DBEAFE"),  # sky-100
    }

    ws["A1"] = "Language Editing Rules"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = brand_fill
    ws.merge_cells("A1:H1")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    meta_rows = [
        ("Project",            code),
        ("Profile",            (cfg.get("profile_key") or "uk").upper()),
        ("Profile name",       cfg.get("name") or ""),
        ("Total rules",        str(len(rules))),
        ("Enabled rules",      f"{enabled_count} / {len(rules)}"),
        ("Variants in dict",   str(len(cfg.get("variant_to_canonical") or {}))),
    ]
    for i, (k, v) in enumerate(meta_rows, start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color="475569")
        ws.cell(row=i, column=2, value=v).font = Font(color="0F172A")

    # ── Rules table ───────────────────────────────────────────────────────
    header_row_idx = len(meta_rows) + 3  # one blank gap
    headers = ["Rule ID", "Message / Description", "Category", "Severity", "Type", "Enabled", "Pattern", "Replacement"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row_idx, column=col, value=h)
        c.font = Font(bold=True, color="312E81")
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")
        c.border = border

    for i, r in enumerate(rules, start=header_row_idx + 1):
        enabled = r.get("enabled", True)
        row_fill = enabled_fill if enabled else disabled_fill
        sev = (r.get("severity") or "").lower()
        values = [
            r.get("id", ""),
            r.get("message", ""),
            (r.get("category") or "").title(),
            sev.title(),
            r.get("type", ""),
            "YES" if enabled else "no",
            r.get("pattern", "") or "",
            "" if r.get("replacement") is None else str(r.get("replacement")),
        ]
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=val)
            c.fill = row_fill
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=col == 2)
            if col == 1:
                c.font = Font(name="Consolas", bold=True, color="0F172A" if enabled else "94A3B8")
            elif col == 4 and sev in sev_fills:
                c.fill = sev_fills[sev]
                c.font = Font(bold=True, color="0F172A" if enabled else "94A3B8")
            elif col == 6:
                c.font = Font(bold=True, color="059669" if enabled else "94A3B8")
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col in (7, 8):
                c.font = Font(name="Consolas", color="334155" if enabled else "94A3B8")
            else:
                c.font = Font(color="0F172A" if enabled else "94A3B8")

    # Column widths
    widths = [14, 50, 14, 14, 10, 10, 36, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze header
    ws.freeze_panes = ws.cell(row=header_row_idx + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    headers = {"Content-Disposition": f'attachment; filename="{code}_language_rules.xlsx"'}
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.post("/files/{file_id}/language-edit/analyze")
def start_language_edit_analysis(
    file_id: int,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Triggers rule engine analysis on target manuscript file using active CE support language rules."""
    _require_language_user(current_user)
    file_record = db.query(models.File).filter(models.File.id == file_id).first()
    if not file_record:
        raise HTTPException(status_code=404, detail=f"File ID {file_id} not found.")

    if not file_record.path or not os.path.exists(file_record.path):
        raise HTTPException(status_code=400, detail="Target file does not exist on disk.")

    if not file_record.filename.endswith(".docx"):
        raise HTTPException(status_code=400, detail="Language Editing is only supported for .docx files.")

    # Load project language rules from CE support
    project_rules_cfg = get_project_language_rules(db, project_id=file_record.project_id)
    rules = load_rules_from_dict(project_rules_cfg.get("rules", []))
    dictionary = load_house_style_from_dict(project_rules_cfg.get("variant_to_canonical", {}))

    # Read paragraphs from DOCX
    try:
        _, paragraphs = docx_io.read_paragraphs(file_record.path)
    except Exception as exc:
        logger.error("Failed to read DOCX at %s: %s", file_record.path, exc)
        raise HTTPException(status_code=500, detail="Failed to read target DOCX file.")

    job_uuid = datetime.datetime.utcnow().strftime("%Y%m%d%H%M%S_") + str(uuid.uuid4())[:8]
    job = LanguageEditJob(
        job_id=job_uuid,
        file_id=file_record.id,
        project_id=file_record.project_id,
        status="in_review",
        total_findings=0,
    )
    db.add(job)
    db.commit()

    total_findings_count = 0
    findings_to_insert = []

    for para_idx, text in paragraphs:
        sents = segmenter.sentences(text)
        found = engine.analyze(text, rules, dictionary, sents)
        for fd in found:
            total_findings_count += 1
            finding_row = LanguageEditFinding(
                job_id=job_uuid,
                file_id=file_record.id,
                para_index=para_idx,
                start_offset=fd.start,
                end_offset=fd.end,
                rule_id=fd.rule_id,
                category=fd.category,
                severity=fd.severity,
                original_text=fd.original,
                suggestion=fd.suggestion,
                message=fd.message,
                autofixable=fd.autofixable,
                status="pending",
            )
            findings_to_insert.append(finding_row)

    if findings_to_insert:
        db.bulk_save_objects(findings_to_insert)

    job.total_findings = total_findings_count
    db.commit()

    return {
        "ok": True,
        "job_id": job_uuid,
        "total_findings": total_findings_count,
        "file_name": file_record.filename,
    }


@router.get("/language-edit/jobs/{job_id}/findings")
def get_job_findings(
    job_id: str,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Returns findings and statistics for a given Language Edit job."""
    _require_language_user(current_user)
    job = db.query(LanguageEditJob).filter(LanguageEditJob.job_id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail=f"Job {job_id} not found.")

    findings = (
        db.query(LanguageEditFinding)
        .filter(LanguageEditFinding.job_id == job_id)
        .order_by(LanguageEditFinding.para_index, LanguageEditFinding.start_offset)
        .all()
    )

    return {
        "job": {
            "job_id": job.job_id,
            "file_id": job.file_id,
            "project_id": job.project_id,
            "status": job.status,
            "total_findings": job.total_findings,
            "accepted_count": job.accepted_count,
            "edited_count": job.edited_count,
            "rejected_count": job.rejected_count,
        },
        "findings": [
            {
                "id": f.id,
                "para_index": f.para_index,
                "start_offset": f.start_offset,
                "end_offset": f.end_offset,
                "rule_id": f.rule_id,
                "category": f.category,
                "severity": f.severity,
                "original_text": f.original_text,
                "suggestion": f.suggestion,
                "message": f.message,
                "autofixable": f.autofixable,
                "status": f.status,
                "edited_text": f.edited_text,
                "reviewer": f.reviewer,
                "decided_at": f.decided_at.isoformat() if f.decided_at else None,
            }
            for f in findings
        ]
    }


@router.post("/language-edit/findings/{finding_id}/decide")
def decide_finding(
    finding_id: int,
    body: DecisionRequest,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Updates editor decision (accepted, edited, rejected, highlighted) for a finding."""
    _require_language_user(current_user)
    finding = db.query(LanguageEditFinding).filter(LanguageEditFinding.id == finding_id).first()
    if not finding:
        raise HTTPException(status_code=404, detail=f"Finding ID {finding_id} not found.")

    if body.status not in ("accepted", "edited", "rejected", "highlighted"):
        raise HTTPException(status_code=400, detail="Invalid status. Must be accepted, edited, rejected, or highlighted.")

    old_status = finding.status
    finding.status = body.status
    finding.edited_text = body.edited_text if body.status == "edited" else None
    finding.reviewer = body.reviewer or "editor"
    finding.decided_at = datetime.datetime.utcnow()

    # Update counters on parent job
    job = db.query(LanguageEditJob).filter(LanguageEditJob.job_id == finding.job_id).first()
    if job:
        # Decrement old counter if replacing decision
        if old_status == "accepted":
            job.accepted_count = max(0, job.accepted_count - 1)
        elif old_status == "edited":
            job.edited_count = max(0, job.edited_count - 1)
        elif old_status == "rejected":
            job.rejected_count = max(0, job.rejected_count - 1)

        # Increment new counter
        if body.status == "accepted":
            job.accepted_count += 1
        elif body.status == "edited":
            job.edited_count += 1
        elif body.status == "rejected":
            job.rejected_count += 1

    db.commit()
    return {"ok": True, "finding_id": finding.id, "status": finding.status}


@router.post("/language-edit/jobs/{job_id}/export")
def export_redlined_docx(
    job_id: str,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """Exports redlined DOCX with Word Tracked Changes and XML highlights, saving it in the same location as a new version."""
    _require_language_user(current_user)
    job = db.query(LanguageEditJob).filter(LanguageEditJob.job_id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail=f"Job {job_id} not found.")

    file_record = db.query(models.File).filter(models.File.id == job.file_id).first()
    if not file_record or not os.path.exists(file_record.path):
        raise HTTPException(status_code=404, detail="Source DOCX file not found.")

    active_findings = (
        db.query(LanguageEditFinding)
        .filter(
            LanguageEditFinding.job_id == job_id,
            LanguageEditFinding.status.in_(["accepted", "edited", "highlighted"])
        )
        .order_by(LanguageEditFinding.para_index, LanguageEditFinding.start_offset.desc())
        .all()
    )

    doc, _ = docx_io.read_paragraphs(file_record.path)

    # Apply right-to-left within each paragraph so offsets remain valid
    for f in active_findings:
        if 0 <= f.para_index < len(doc.paragraphs):
            para = doc.paragraphs[f.para_index]
            if f.status == "highlighted":
                color = "cyan" if f.category == "spelling" else "magenta" if f.category == "sentence" else "yellow"
                docx_io.apply_highlight_change(para, start=f.start_offset, end=f.end_offset, color=color)
            else:
                target_text = f.edited_text if f.status == "edited" and f.edited_text else f.suggestion
                docx_io.apply_tracked_change(
                    para,
                    start=f.start_offset,
                    end=f.end_offset,
                    new_text=target_text,
                    author="LangQA"
                )

    # 1. Archive previous file version in Archive/ folder before overwriting
    from app.domains.files import version_service
    base_path = os.path.dirname(file_record.path)
    try:
        version_service.archive_existing_file(
            db,
            existing_file=file_record,
            base_path=base_path,
            uploaded_by_id=None,
            reason="Language Edit review completed"
        )
    except Exception as err:
        logger.warning(f"Failed to archive previous file version: {err}")

    # 2. Save redlined DOCX to the exact same location as the new file version
    doc.save(file_record.path)

    # 3. Increment file record version and update timestamp
    file_record.version = (file_record.version or 1) + 1
    file_record.updated_at = datetime.datetime.utcnow()

    # Save export copy for download response
    out_dir = os.path.join(base_path, "jobs")
    os.makedirs(out_dir, exist_ok=True)
    filename_stem = os.path.splitext(file_record.filename)[0]
    out_path = os.path.join(out_dir, f"{filename_stem}_redline.docx")
    doc.save(out_path)

    job.status = "completed"
    db.commit()

    headers = {
        "X-File-Id": str(file_record.id),
        "X-File-Version": str(file_record.version),
        "Access-Control-Expose-Headers": "X-File-Id, X-File-Version"
    }

    return FileResponse(
        out_path,
        filename=f"{filename_stem}_redline.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers
    )



@router.get("/medical/umls-search")
def search_medical_term_umls(
    term: str,
    api_key: Optional[str] = None,
    current_user: Optional[User] = Depends(get_current_user_from_cookie),
):
    """
    Search medical and pharmaceutical terms using NLM UMLS Terminology Services API.
    Returns matched concepts, canonical names, and CUIs.
    """
    _require_language_user(current_user)
    if not term or not term.strip():
        raise HTTPException(status_code=400, detail="Query term parameter is required.")

    from app.services.umls_service import validate_medical_term
    result = validate_medical_term(term, api_key=api_key)
    return {"ok": True, "data": result}

