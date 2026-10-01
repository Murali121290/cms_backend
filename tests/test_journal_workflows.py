import io

import pytest
from docx import Document

from tests.test_journal_stage1_checks import build_manuscript


@pytest.fixture()
def api(auth_cookie_client, admin_user, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return auth_cookie_client(admin_user)


@pytest.fixture()
def client_row(api):
    return api.post("/api/v2/journals/clients", json={"client_code": "ELSA-01", "publisher_name": "Elsevier"}).json()


def docx_bytes(title, doi=None):
    doc = Document()
    doc.add_paragraph(title, style="Title")
    if doi:
        doc.add_paragraph(f"DOI: {doi}")
    doc.add_paragraph("Abstract Short abstract.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def finish_step(api, aid, step):
    """Mark the step's open errors fixed, then finish it (accepting warnings)."""
    module = step
    for i in api.get(f"/api/v2/journals/articles/{aid}/issues", params={"module": module, "issue_status": "open"}).json():
        if i["severity"] == "error":
            api.patch(f"/api/v2/journals/articles/{aid}/issues/{i['id']}", json={"action": "accept"})
    res = api.post(f"/api/v2/journals/articles/{aid}/pre-editing/{step}/finish", json={"accept_warnings": True})
    assert res.status_code == 200, res.text
    return res.json()


def test_default_workflows_are_listed(api):
    wfs = api.get("/api/v2/journals/workflows").json()
    assert [w["name"] for w in wfs][0] == "Full production"
    xml_only = next(w for w in wfs if w["name"].startswith("XML only"))
    assert xml_only["stage_numbers"] == [1, 2, 3, 7]
    assert xml_only["stages"][-1] == "7. Final Delivery"


def test_create_workflow_validates_stages(api):
    assert api.post("/api/v2/journals/workflows", json={"name": "Bad", "stage_numbers": [0, 9]}).status_code == 400
    ok = api.post("/api/v2/journals/workflows", json={"name": "Editing only", "stage_numbers": [3, 1, 2, 2, 7]})
    assert ok.status_code == 201 and ok.json()["stage_numbers"] == [1, 2, 3, 7]
    assert api.post("/api/v2/journals/workflows", json={"name": "Old eight", "stage_numbers": [1, 8]}).status_code == 400


def test_journal_gets_default_workflow_and_overview(api, client_row):
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JAIS", "journal_title": "JAIS"}).json()
    assert j["workflow_name"] == "Full production"
    assert api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "X", "journal_title": "X", "workflow_id": 999}).status_code == 400
    assert api.post("/api/v2/journals", json={"client_id": 999, "journal_code": "Y", "journal_title": "Y"}).status_code == 404

    detail = api.get(f"/api/v2/journals/{j['id']}").json()
    assert detail["client_code"] == "ELSA-01" and len(detail["stages"]) == 7
    assert api.get("/api/v2/journals/overview", params={"client_id": client_row["id"]}).json()[0]["journal_code"] == "JAIS"
    overview = api.get("/api/v2/journals/clients/overview").json()
    assert overview[0]["journal_count"] == 1 and overview[0]["articles"]["total"] == 0
    assert api.get(f"/api/v2/journals/clients/{client_row['id']}").json()["publisher_name"] == "Elsevier"


def test_batch_upload_creates_articles_on_journal_workflow(api, client_row, tmp_path):
    wf = next(w for w in api.get("/api/v2/journals/workflows").json() if w["name"].startswith("XML only"))
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JX", "journal_title": "JX", "workflow_id": wf["id"]}).json()
    manuscript = build_manuscript(tmp_path / "m.docx").read_bytes()
    files = [
        ("files", ("first.docx", docx_bytes("First article", "10.1016/j.jx.2026.001"), "application/octet-stream")),
        ("files", ("dup.docx", docx_bytes("Duplicate", "10.1016/j.jx.2026.001"), "application/octet-stream")),
        ("files", ("notes.txt", b"hello", "text/plain")),
        ("files", ("second.docx", manuscript, "application/octet-stream")),
    ]
    res = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=files)
    assert res.status_code == 200, res.text
    body = res.json()
    assert [c["filename"] for c in body["created"]] == ["first.docx", "second.docx"]
    assert {f["filename"] for f in body["failed"]} == {"dup.docx", "notes.txt"}
    assert "already exists" in next(f for f in body["failed"] if f["filename"] == "dup.docx")["error"]

    rows = api.get(f"/api/v2/journals/{j['id']}/articles").json()
    assert len(rows) == 2
    assert [s["stage_number"] for s in rows[0]["stages"]] == [1, 2, 3, 7]
    assert rows[0]["current_stage"] == "1. Pre-Editing"
    assert rows[0]["stages"][0]["stage_status"] == "In-progress"
    files_of_first = api.get(f"/api/v2/journals/articles/{rows[0]['id']}/files").json()
    assert files_of_first[0]["category"] == "Manuscript"


def test_advance_follows_the_workflow_and_skips_unselected_stages(api, client_row, db_session):
    from app.domains.journals.models import JournalArticle
    wf = api.post("/api/v2/journals/workflows", json={"name": "Two step", "stage_numbers": [2, 7]}).json()
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "J2", "journal_title": "J2", "workflow_id": wf["id"]}).json()
    art = api.post("/api/v2/journals/articles", json={"journal_id": j["id"], "article_title": "A"}).json()
    assert art["current_stage"] == "2. Language Editing"

    res = api.post(f"/api/v2/journals/articles/{art['id']}/advance-stage", json={})
    assert res.status_code == 200 and res.json()["new_stage"] == "7. Final Delivery"
    done = api.post(f"/api/v2/journals/articles/{art['id']}/advance-stage", json={}).json()
    assert done["status"] == "completed"
    db_session.expire_all()
    assert db_session.get(JournalArticle, art["id"]).status == "Completed"
    row = api.get(f"/api/v2/journals/{j['id']}/articles").json()[0]
    assert [s["stage_status"] for s in row["stages"]] == ["Completed", "Completed"]
    assert api.get("/api/v2/journals/clients/overview").json()[0]["articles"]["completed"] == 1


def test_workspace_reads_state_without_reprocessing(api, client_row, tmp_path, monkeypatch):
    from app.domains.journals.checks import references as references_check
    monkeypatch.setattr(references_check, "crossref_lookup", lambda title, year: [])
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JW", "journal_title": "JW"}).json()
    files = [("files", ("m.docx", build_manuscript(tmp_path / "m.docx").read_bytes(), "application/octet-stream"))]
    aid = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=files).json()["created"][0]["id"]

    ws = api.get(f"/api/v2/journals/articles/{aid}/workspace").json()
    assert ws["xhtml"] is None and ws["check_runs"] == {} and len(ws["stages"]) == 7
    assert ws["journal"]["journal_code"] == "JW"

    api.post(f"/api/v2/journals/articles/{aid}/process-pre-editing")
    api.post(f"/api/v2/journals/articles/{aid}/checks/references/run")
    ws = api.get(f"/api/v2/journals/articles/{aid}/workspace").json()
    assert ws["xhtml"]["file"]["version"] == 1 and 'data-para-idx' in ws["xhtml"]["content"]
    assert ws["check_runs"]["structuring"]["status"] == "Completed"
    assert ws["open_issues"]["references"]["error"] >= 1
    # Reading the workspace again does not create a new XHTML version.
    api.get(f"/api/v2/journals/articles/{aid}/workspace")
    xhtml_files = [f for f in api.get(f"/api/v2/journals/articles/{aid}/files").json() if f["category"] == "XHTML"]
    assert len(xhtml_files) == 1


def test_zip_upload_rejects_duplicate_doi(api, client_row, tmp_path):
    import zipfile
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JZ", "journal_title": "JZ"}).json()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("manuscript.docx", docx_bytes("Zip article", "10.1016/j.jz.2026.001"))
    data = buf.getvalue()
    first = api.post("/api/v2/journals/articles/upload-zip", data={"journal_id": j["id"]}, files={"file": ("a.zip", data, "application/zip")})
    assert first.status_code == 200, first.text
    again = api.post("/api/v2/journals/articles/upload-zip", data={"journal_id": j["id"]}, files={"file": ("b.zip", data, "application/zip")})
    assert again.status_code == 409


def test_parenthesis_citations_are_recognised(api, client_row, tmp_path, monkeypatch):
    from app.domains.journals.checks import references as references_check
    monkeypatch.setattr(references_check, "crossref_lookup", lambda title, year: [])
    doc = Document()
    doc.add_paragraph("Cirrhosis is common (1). Early stages are asymptomatic (2, 3), and aged 35-60 (4). Rates rose in 2019 (2019).")
    doc.add_paragraph("Also see (9).")
    doc.add_paragraph("References")
    for n in range(1, 5):
        doc.add_paragraph(f"{n}. Author A. Title {n}. J Test. 2020;1:1-2.")
    path = tmp_path / "paren.docx"
    doc.save(path)
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JP", "journal_title": "JP"}).json()
    aid = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=[("files", ("paren.docx", path.read_bytes(), "application/octet-stream"))]).json()["created"][0]["id"]
    api.post(f"/api/v2/journals/articles/{aid}/checks/references/run")
    issues = [i for i in api.get(f"/api/v2/journals/articles/{aid}/issues", params={"module": "references"}).json() if i["status"] == "open"]
    assert [i["title"] for i in issues if i["rule_id"] == "REF-X01"] == ["Citation (9) has no matching reference"]
    assert not [i for i in issues if i["rule_id"] == "REF-C01"]  # all four references are cited
    x01 = next(i for i in issues if i["rule_id"] == "REF-X01")
    assert x01["location"]["surface"] == "(9)"


def test_saving_editor_html_updates_edited_docx_and_reruns_checks(api, client_row, tmp_path, monkeypatch, db_session):
    import re
    from app.domains.journals.checks import references as references_check
    from app.domains.journals.models import JournalArticle
    monkeypatch.setattr(references_check, "crossref_lookup", lambda title, year: [])
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JE", "journal_title": "JE"}).json()
    files = [("files", ("m.docx", build_manuscript(tmp_path / "m.docx").read_bytes(), "application/octet-stream"))]
    aid = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=files).json()["created"][0]["id"]
    pre = api.post(f"/api/v2/journals/articles/{aid}/process-pre-editing").json()
    open_rules = lambda: {i["rule_id"] for i in api.get(f"/api/v2/journals/articles/{aid}/issues", params={"module": "structuring"}).json() if i["status"] == "open"}
    # structuring_lib tagged the numbered headings, so none is left as body text.
    assert pre["structuring"]["paragraphs"] > 0 and pre["structuring"]["front_matter"]["abstract"]
    assert "STR-H02" not in open_rules()

    html = api.get(f"/api/v2/journals/articles/{aid}/workspace").json()["xhtml"]["content"]
    # Retag "1 Introduction" as body text (TXT), as the editor's PARA panel would.
    html2, n = re.subn(r'<(p|h\d) class="[^"]*" data-style-label="[^"]*"([^>]*)>((?:<span[^>]*>)?)1 Introduction',
                       r'<p class="TXT" data-style-label="TXT"\g<2>>\g<3>1 Introduction', html, count=1)
    assert n == 1
    res = api.put(f"/api/v2/journals/articles/{aid}/xhtml", json={"html_content": html2})
    assert res.status_code == 200, res.text
    assert res.json()["xhtml_version"] == 2
    assert "STR-H02" in open_rules()  # the saved edit reached the manuscript the checks read

    db_session.expire_all()
    art = db_session.get(JournalArticle, aid)
    assert art.edited_docx_path.endswith("_structured.docx") and art.edited_docx_path != art.original_docx_path
    from docx import Document as D
    styles = {p.text: p.style.name for p in D(art.edited_docx_path).paragraphs}
    assert styles["1 Introduction"] == "TXT"
    assert styles["A Benchmark for Causal Reasoning"] == "ArticleTitle"
    assert D(art.original_docx_path).paragraphs[2].style.name == "Heading 1"  # the upload is untouched


def test_front_matter_after_keywords_is_not_left_as_headings(tmp_path):
    from app.domains.journals.structuring import structure_manuscript
    doc = Document()
    for t in ["Original article",
              "Nutritional status and respiratory muscle strength in sarcoidosis patients",
              "Şeyma Tunç1, Pınar Yıldız2, Mehmet Ali Sungur3",
              "1 Department of Physiotherapy, Kastamonu University, Kastamonu, Turkey",
              "Abstract. Background and aim: Sarcoidosis is a systemic granulomatous disease. Methods: 60 patients. Results: lower strength.",
              "Key words: sarcoidosis, nutritional status, respiratory muscle strength, fatigue",
              "Received: 12 May 2025",
              "Accepted: 5 August 2025",
              "Correspondence: Şeyma Tunç",
              "Kastamonu, 37100 Turkey",
              "E-mail: seyma@example.edu",
              "ORCID: 0000-0001-7452-1173",
              "Introduction",
              "Sarcoidosis is a multisystem disease of unknown cause (1).",
              "References",
              "1. Grunewald J. Sarcoidosis. Nat Rev Dis Primers. 2019;5:45."]:
        doc.add_paragraph(t)
    src, out = tmp_path / "s.docx", tmp_path / "o.docx"
    doc.save(src)
    info = structure_manuscript(str(src), str(out))
    paras = Document(out).paragraphs
    style = lambda start: next(p.style.name for p in paras if p.text.startswith(start))  # noqa: E731
    assert style("Original article") == "ArticleType"
    assert style("Nutritional status and") == "ArticleTitle"
    assert style("Şeyma Tunç1") == "ArticleAuthor"
    assert style("Abstract. Background") == "Abstract"
    assert style("Key words:") == "Keywords"
    for line in ("Received:", "Accepted:", "Correspondence:", "E-mail:", "ORCID:"):
        assert style(line) == "PMI", line
    assert style("Kastamonu, 37100") == "ArticleAffiliation"
    styles = {"Introduction": style("Introduction")}
    assert styles["Introduction"] == "H1"
    assert info["front_matter"]["title"].startswith("Nutritional status")


def test_technical_and_language_checks(api, client_row, tmp_path, monkeypatch):
    from app.domains.journals.checks import references as references_check
    monkeypatch.setattr(references_check, "crossref_lookup", lambda title, year: [])
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JT", "journal_title": "JT"}).json()
    api.post(f"/api/v2/journals/{j['id']}/stylesheets", json={"name": "S", "style_rules": {"keywords": {"min": 4, "max": 6}}})
    api.post(f"/api/v2/journals/{j['id']}/grammarsheets", json={"name": "G", "language_variant": "US_English",
             "grammar_rules": {"terms": [{"find": "e-mail", "replace": "email"}]}})
    doc = Document()
    for t in ["A study of things", "Abstract We looked at things.", "Key words: things, stuff",
              "Introduction", "Samples took 10ms and 27 percent failed across ages 10-20 (Fig. 1).",
              "In order to utilise the data, which is noisy, we sent an e-mail to cells, tissues and organs.",
              "Figure 1. Overview.", "Figure 2. Never cited.", "References", "1. A B. T. J. 2020;1:1."]:
        doc.add_paragraph(t)
    path = tmp_path / "t.docx"
    doc.save(path)
    aid = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=[("files", ("t.docx", path.read_bytes(), "application/octet-stream"))]).json()["created"][0]["id"]

    assert api.post(f"/api/v2/journals/articles/{aid}/checks/technical/run").status_code == 200
    tech = {i["rule_id"]: i for i in api.get(f"/api/v2/journals/articles/{aid}/issues", params={"module": "technical", "issue_status": "open"}).json()}
    assert tech["JT-FIG"]["suggestion"]["to"] == "Figure 1"
    assert tech["JT-UNIT"]["suggestion"]["to"] == "10 ms"
    assert tech["JT-PCT"]["suggestion"]["to"] == "27%"
    assert tech["JT-RANGE"]["suggestion"]["to"] == "10–20"
    assert tech["JT-CITE"]["title"] == "Figure 2 is never mentioned in the text"
    assert tech["JT-KWD"]["title"].startswith("2 keywords")

    assert api.post(f"/api/v2/journals/articles/{aid}/checks/language/run").status_code == 200
    lang = api.get(f"/api/v2/journals/articles/{aid}/issues", params={"module": "language", "issue_status": "open"}).json()
    subs = {(i["suggestion"] or {}).get("from"): (i["suggestion"] or {}).get("to") for i in lang}
    assert subs.get("In order to") == "To"
    assert subs.get("utilise") == "utilize"
    assert subs.get("e-mail") == "email"
    assert any(i["rule_id"] == "JL-SERIAL" for i in lang)
    assert all(i["severity"] != "error" for i in lang)  # language findings never block a stage


def test_local_reference_processing_styles_refs_and_builds_element_citations(api, client_row, tmp_path, monkeypatch, db_session):
    from app.domains.journals.checks import references as references_check
    from app.domains.journals.models import JournalArticle
    from app.domains.journals.jats.validator import validate_jats
    monkeypatch.setattr(references_check, "crossref_lookup", lambda title, year: [])
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JR", "journal_title": "JR", "issn_print": "1234-5678"}).json()
    doc = Document()
    for t in ["Nutrition in sarcoidosis patients", "Abstract Short.", "Key words: a, b", "Introduction",
              "Sarcoidosis is systemic (1). Strength fell (2).", "References",
              "1. Grunewald J, Grutters JC, Arkema EV, et al. Sarcoidosis. Nat Rev Dis Primers. 2019;5(1):45-60. doi:10.1038/s41572-019-0096-x",
              "2. Pearl J, Mackenzie D. The Book of Why. New York: Basic Books; 2018."]:
        doc.add_paragraph(t)
    path = tmp_path / "r.docx"
    doc.save(path)
    aid = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=[("files", ("r.docx", path.read_bytes(), "application/octet-stream"))]).json()["created"][0]["id"]
    api.post(f"/api/v2/journals/articles/{aid}/process-pre-editing")
    finish_step(api, aid, "structuring")
    # The References step styles the references (bib_*) and the citations (cite_bib).
    ref = api.post(f"/api/v2/journals/articles/{aid}/pre-editing/references/run")
    assert ref.status_code == 200, ref.text
    styling = ref.json()["reference_styling"]
    assert styling["references_styled"] == 2 and styling["citations_styled"] == 2, styling
    assert "cite_bib" in api.get(f"/api/v2/journals/articles/{aid}/workspace").json()["xhtml"]["content"]

    # "Process references" can still be run again; styled references are left as they are.
    res = api.post(f"/api/v2/journals/articles/{aid}/references/process")
    assert res.status_code == 202 and res.json()["engine"] == "local"
    status = api.get(f"/api/v2/journals/articles/{aid}/references").json()
    assert status["status"].startswith("completed (local: 0 references structured"), status
    assert status["reports"], "the local log is kept as a report file"

    db_session.expire_all()
    art = db_session.get(JournalArticle, aid)
    styles = {}
    for p in Document(art.edited_docx_path).paragraphs:
        for r in p.runs:
            if r.style is not None and r.style.name.startswith("bib_"):
                styles.setdefault(r.style.name, []).append(r.text)
    assert styles["bib_surname"][:3] == ["Grunewald", "Grutters", "Arkema"]
    assert styles["bib_journal"] == ["Nat Rev Dis Primers"] and styles["bib_doi"] == ["10.1038/s41572-019-0096-x"]
    assert styles["bib_book"] == ["The Book of Why"] and styles["bib_publisher"] == ["Basic Books"]
    # the text itself is unchanged
    assert any(p.text.startswith("1. Grunewald J, Grutters JC") for p in Document(art.edited_docx_path).paragraphs)
    snapshots = [f for f in api.get(f"/api/v2/journals/articles/{aid}/files").json() if f["category"] == "Working_Copy"]
    assert len(snapshots) == 1  # the pre-reference working copy is kept as a version
    structuring = {i["rule_id"] for i in api.get(f"/api/v2/journals/articles/{aid}/issues", params={"module": "structuring", "issue_status": "open"}).json()}
    assert "STR-CS-09" not in structuring  # bib_* styles are not "unmapped"

    art.current_stage = "3. XML Conversion"
    db_session.commit()
    conv = api.post(f"/api/v2/journals/articles/{aid}/xml/convert").json()
    xml = api.get(f"/api/v2/journals/articles/{aid}/xml").json()["content"]
    assert '<element-citation publication-type="journal">' in xml
    assert "<surname>Grunewald</surname>" in xml and "<source>Nat Rev Dis Primers</source>" in xml
    assert '<pub-id pub-id-type="doi">10.1038/s41572-019-0096-x</pub-id>' in xml
    assert '<element-citation publication-type="book">' in xml and "<publisher-name>Basic Books</publisher-name>" in xml
    assert [f.rule_id for f in validate_jats(xml.encode("utf-8")) if f.severity == "error"] == [], conv


def test_book_review_shadow_file(api, client_row, tmp_path, monkeypatch, db_session):
    import os
    from app.domains.journals.checks import references as references_check
    from app.domains.journals.models import JournalArticle
    monkeypatch.setattr(references_check, "crossref_lookup", lambda title, year: [])
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JB", "journal_title": "JB"}).json()
    files = [("files", ("m.docx", build_manuscript(tmp_path / "m.docx").read_bytes(), "application/octet-stream"))]
    aid = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=files).json()["created"][0]["id"]
    assert api.post(f"/api/v2/journals/articles/{aid}/review-file").status_code == 400  # no working copy yet

    api.post(f"/api/v2/journals/articles/{aid}/process-pre-editing")
    first = api.post(f"/api/v2/journals/articles/{aid}/review-file").json()
    again = api.post(f"/api/v2/journals/articles/{aid}/review-file").json()
    assert first == again  # one shadow file per article

    from app import models
    from app.domains.projects.models import Project
    db_session.expire_all()
    art = db_session.get(JournalArticle, aid)
    f = db_session.get(models.File, first["file_id"])
    assert os.path.abspath(art.edited_docx_path) == f.path and f.chapter_id is None
    p = db_session.get(Project, first["project_id"])
    assert (p.project_code, p.is_deleted) == ("JRNL-JB", True)

    xhtml = api.get(f"/api/v2/files/{first['file_id']}/xhtml-runs")  # the book structuring page's loader
    assert xhtml.status_code == 200 and "data-para-idx" in xhtml.json()["content"]
    listed = api.get("/api/v2/projects").json()
    codes = [x.get("code") or x.get("project_code") for x in (listed.get("projects") if isinstance(listed, dict) else listed)]
    assert "JRNL-JB" not in codes

    assert api.get(f"/api/v2/journals/articles/{aid}/workspace").json()["working_copy_changed"] is False
    import time
    time.sleep(1.2)
    from docx import Document as D
    d = D(f.path)
    d.add_paragraph("Added in the book technical review.")
    d.save(f.path)
    os.utime(f.path, (time.time() + 5, time.time() + 5))
    assert api.get(f"/api/v2/journals/articles/{aid}/workspace").json()["working_copy_changed"] is True
    api.post(f"/api/v2/journals/articles/{aid}/process-pre-editing")  # "Refresh": re-converts the working copy, keeps the edit
    ws = api.get(f"/api/v2/journals/articles/{aid}/workspace").json()
    assert "Added in the book technical review." in ws["xhtml"]["content"]


def test_ia_rules_become_the_hidden_projects_active_stylesheet(api, client_row):
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JI", "journal_title": "JI"}).json()
    data = api.get(f"/api/v2/journals/{j['id']}/ia-rules").json()
    assert len(data["catalog"]) > 50 and data["selected"] == [] and data["stylesheet"] is None
    pick = [r for r in data["catalog"] if r["element"] in ("Figure", "Ranges")][:4]
    bogus = {"element": "Nope", "subtype": "x", "pattern": "y"}
    saved = api.put(f"/api/v2/journals/{j['id']}/ia-rules", json={"selected_ia_rows": pick + [pick[0], bogus]}).json()
    assert len(saved["selected"]) == 4  # duplicates and unknown rows dropped
    book = api.get(f"/api/v2/projects/{saved['project_id']}/stylesheets").json()
    active = book["active_stylesheet"]
    assert active and active["id"] == saved["stylesheet"]["id"]
    assert {(r["element"], r["pattern"]) for r in active["selected_ia_rows"]} == {(r["element"], r["pattern"]) for r in pick}
    again = api.put(f"/api/v2/journals/{j['id']}/ia-rules", json={"selected_ia_rows": pick[:1]}).json()
    assert again["stylesheet"]["id"] == saved["stylesheet"]["id"]  # same stylesheet updated, not a new one
    assert len(api.get(f"/api/v2/journals/{j['id']}/ia-rules").json()["selected"]) == 1


def test_delete_article_removes_rows_files_and_shadow_file(api, client_row, tmp_path, monkeypatch, db_session):
    import os
    from app import models
    from app.domains.journals.checks import references as references_check
    from app.domains.journals.models import JournalArticle, JournalIssue, JournalStageDetail
    monkeypatch.setattr(references_check, "crossref_lookup", lambda title, year: [])
    j = api.post("/api/v2/journals", json={"client_id": client_row["id"], "journal_code": "JD", "journal_title": "JD"}).json()
    files = [("files", ("m.docx", build_manuscript(tmp_path / "m.docx").read_bytes(), "application/octet-stream"))]
    aid = api.post(f"/api/v2/journals/{j['id']}/articles/upload", files=files).json()["created"][0]["id"]
    api.post(f"/api/v2/journals/articles/{aid}/process-pre-editing")
    shadow = api.post(f"/api/v2/journals/articles/{aid}/review-file").json()["file_id"]
    db_session.expire_all()
    art = db_session.get(JournalArticle, aid)
    paths = [art.original_docx_path, art.edited_docx_path, art.xhtml_path]
    assert all(os.path.exists(p) for p in paths)

    res = api.delete(f"/api/v2/journals/articles/{aid}")
    assert res.status_code == 200, res.text
    db_session.expire_all()
    assert db_session.get(JournalArticle, aid) is None
    assert db_session.query(JournalIssue).filter_by(article_id=aid).count() == 0
    assert db_session.query(JournalStageDetail).filter_by(article_id=aid).count() == 0
    assert db_session.get(models.File, shadow) is None
    assert not any(os.path.exists(p) for p in paths)
    assert api.get(f"/api/v2/journals/{j['id']}/articles").json() == []
    assert api.delete(f"/api/v2/journals/articles/{aid}").status_code == 404
