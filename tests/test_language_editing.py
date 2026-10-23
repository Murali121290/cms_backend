"""
Unit tests for Ninja Inkflow Language Editing module.
"""
import pytest
import os
from docx import Document
from app.processing.language_editing import segmenter, engine, rules, docx_io, custom_dict, spell_check


def test_segmenter_abbreviations_and_decimals():
    text = "The results in Fig. 7 were analyzed carefully (3.14). Dr. Smith et al. met the deadline."
    sents = segmenter.sentences(text)
    assert len(sents) == 2
    assert "Fig. 7" in sents[0][2]
    assert "Dr. Smith et al. met" in sents[1][2]


def test_engine_grammar_and_spelling_rules():
    rule_dicts = [
        {"id": "GP001", "category": "grammar", "type": "regex", "pattern": " {2,}", "replacement": " ", "message": "Double space", "severity": "error", "enabled": True},
        {"id": "GP002", "category": "grammar", "type": "regex", "pattern": "\\s+([,;:.!?])", "replacement": "\\1", "message": "Space before punctuation", "severity": "error", "enabled": True},
        {"id": "SP000", "category": "spelling", "type": "dictionary", "message": "House style", "severity": "suggestion", "enabled": True}
    ]
    compiled_rules = rules.load_rules_from_dict(rule_dicts)
    house_style = rules.load_house_style_from_dict({"color": "colour", "organization": "organisation"})

    text = "The  color pattern was analyzed carefully . The organization met ."
    sents = segmenter.sentences(text)
    findings = engine.analyze(text, compiled_rules, house_style, sents)

    originals = [f.original for f in findings]
    assert "  " in originals or "color" in originals or " ." in originals


def test_tracked_changes_exporter(tmp_path):
    doc = Document()
    doc.add_paragraph("The results were analyzed carefully .The team met the the deadline.")
    path = os.path.join(tmp_path, "test.docx")
    doc.save(path)

    doc_read, paras = docx_io.read_paragraphs(path)
    p = doc_read.paragraphs[0]
    
    # Apply tracked change
    docx_io.apply_tracked_change(p, start=4, end=11, new_text="outcomes", author="LangQA")
    
    out_path = os.path.join(tmp_path, "test_redline.docx")
    doc_read.save(out_path)
    assert os.path.exists(out_path)


def test_medical_pharmaceutical_profile():
    from app.services.language_rule_service import get_available_style_profiles
    profiles = get_available_style_profiles()
    assert "medical_pharmaceutical" in profiles
    med = profiles["medical_pharmaceutical"]
    assert "anaemia" in med["variant_to_canonical"]
    assert med["variant_to_canonical"]["anaemia"] == "anemia"


def test_umls_service():
    from app.services.umls_service import validate_medical_term
    res = validate_medical_term("ibuprofen")
    assert "valid" in res
    assert res["term"] == "ibuprofen"


def test_paragraph_exclusions(tmp_path):
    doc = Document()
    doc.add_paragraph("<front>")
    doc.add_paragraph("<PT>Part II: The Science of Personal Training")
    doc.add_paragraph("<CN>CHAPTER 3")
    doc.add_paragraph("<CT>Anatomy and Kinesiology")
    doc.add_paragraph("<OBJ1>OBJECTIVES")
    doc.add_paragraph("Personal Trainers should be able to:")
    doc.add_paragraph("<body>")
    doc.add_paragraph("<H1-INTRO>INTRODUCTION")
    doc.add_paragraph("A major goal of exercise training is to improve cardiovascular fitness.")
    doc.add_paragraph("<BXM>")
    doc.add_paragraph("Kinesiology is the study of human movement.")
    doc.add_paragraph("</BXM>")
    doc.add_paragraph("<FIG3.4>")
    doc.add_paragraph("<TAB3.2>")
    doc.add_paragraph("<ref-open>")
    doc.add_paragraph("<REF-N>Smith J. Journal of Science 2020.")
    doc.add_paragraph("<ref-close>")
    
    path = os.path.join(tmp_path, "exclusions_test.docx")
    doc.save(path)

    doc_read, paras = docx_io.read_paragraphs(path)
    filtered_texts = [p_text for idx, p_text in paras]
    
    # Only body paragraphs should be included (<FIG3.4>, <TAB3.2>, <BXM>, </BXM> skipped)
    assert len(filtered_texts) == 2
    assert any("A major goal of exercise training" in t for t in filtered_texts)
    assert any("Kinesiology is the study" in t for t in filtered_texts)
    assert not any("<FIG3.4>" in t for t in filtered_texts)
    assert not any("<TAB3.2>" in t for t in filtered_texts)
    assert not any("</BXM>" in t for t in filtered_texts)


def test_inline_markup_tag_findings_filter():
    from app.processing.language_editing.rules import Rule
    rule = Rule(id="GP002", category="grammar", type="regex", pattern=r"\s+([,;:\.!?])", replacement=r"\1", message="Space before punct", severity="error", enabled=True)
    text = "Kinesiology dates back to the sixth century B.C (10-12-26). </BXM>"
    findings = engine.run_regex_rule(rule, text)
    filtered = engine.filter_markup_tag_findings(text, findings)
    
    # Finding on </BXM> should be filtered out
    assert not any(f.original == " >" or ">" in f.original for f in filtered)


def test_ninja_inkflow_spec_rules():
    from app.processing.language_editing import sentence_checks
    # 1. Test GP005 repeated words guarding that that and had had
    rule_gp005 = rules.Rule(id="GP005", category="grammar", type="regex", pattern=r"\b(\w+)\s+\1\b", replacement=r"\1", message="Repeated word", severity="warning", flags=2, enabled=True).compile()
    text_dup = "He said that that decision was wrong and they had had enough time to met the the deadline."
    findings_dup = engine.run_regex_rule(rule_gp005, text_dup)
    origs = [f.original.lower() for f in findings_dup]
    assert "the the" in origs
    assert "that that" not in origs
    assert "had had" not in origs

    # 2. Test SL003 start_capital display quote guard
    rule_sl003 = rules.Rule(id="SL003", category="sentence", type="function", function="start_capital", severity="warning", enabled=True)
    sent_quote = '"this is a display quote."'
    findings_quote = sentence_checks.start_capital(sent_quote, 0, rule_sl003, {})
    assert len(findings_quote) == 0

    # 3. Test SL004 wordiness phrase replacements
    rule_sl004 = rules.Rule(id="SL004", category="sentence", type="function", function="wordiness", severity="suggestion", enabled=True)
    sent_wordy = "The study was conducted in order to test outcomes due to the fact that data was limited."
    findings_wordy = sentence_checks.wordiness(sent_wordy, 0, rule_sl004, {})
    suggs = [f.suggestion for f in findings_wordy]
    assert "to" in suggs
    assert "because" in suggs


# ---------------------------------------------------------------------------
# Custom dictionary (medical / pharmaceutical) integration
# ---------------------------------------------------------------------------


def _write_utf16_dic(path, lines):
    """Write a .dic file in UTF-16 LE with BOM — matches the shape of client
    deliverables (MedSpl800 was shipped this way)."""
    with open(path, "w", encoding="utf-16", newline="\r\n") as fh:
        fh.write("\n".join(lines) + "\n")


def test_custom_dict_loader_parses_utf16_bom_and_slash(tmp_path):
    dic = tmp_path / "sample.dic"
    _write_utf16_dic(dic, [
        "cholecalciferol",
        "5-hydroxytryptamine",
        "anterior/posterior",
        "Aëtius",
    ])
    words = custom_dict.load_dic_file(dic)
    # Plain long term, both sides of a slash entry, and the alpha portion of
    # a hyphenated chemistry term are all extracted.
    assert "cholecalciferol" in words
    assert "hydroxytryptamine" in words
    assert "anterior" in words and "posterior" in words
    # Diacritic entries are split on non-ASCII letters because the Spell
    # Check tokenizer only recognises [A-Za-z]+ runs — so `Aëtius` in a
    # document also tokenizes to `A`/`tius`. Storing the ASCII fragment
    # (`tius`) is what actually shields the document occurrence.
    assert "tius" in words


def test_custom_dict_loader_skips_metadata_header_blanks_and_comments(tmp_path):
    dic = tmp_path / "with_header.dic"
    _write_utf16_dic(dic, [
        "800",            # Hunspell-style count header — must not become a term
        "# this is a comment",
        "",
        "// another comment",
        "   ",
        "aspirin",
        "ibuprofen",
    ])
    words = custom_dict.load_dic_file(dic)
    assert words == {"aspirin", "ibuprofen"}
    # Specifically prove the metadata/comment lines didn't leak in.
    assert "800" not in words
    assert "comment" not in words


def test_custom_dict_loader_accepts_utf8_without_bom(tmp_path):
    dic = tmp_path / "plain.dic"
    dic.write_text("warfarin\nmetformin\n", encoding="utf-8")
    words = custom_dict.load_dic_file(dic)
    assert words == {"warfarin", "metformin"}


def test_load_all_custom_dicts_honors_env_var(tmp_path, monkeypatch):
    dic = tmp_path / "client_x.dic"
    _write_utf16_dic(dic, ["xyzzyomycin", "foobaritide"])
    monkeypatch.setenv(custom_dict.ENV_VAR, str(dic))
    custom_dict.load_all_custom_dicts.cache_clear()
    try:
        words = custom_dict.load_all_custom_dicts()
    finally:
        custom_dict.load_all_custom_dicts.cache_clear()
    assert "xyzzyomycin" in words
    assert "foobaritide" in words


def test_spell_check_whitelists_medical_dic_term():
    """Long, lowercase medical term from a loaded dic MUST NOT be flagged."""
    sc = spell_check.get_spell_checker(custom_words={"cholecalciferol"})
    if sc is None:
        pytest.skip("pyspellchecker not installed")
    whitelist = spell_check.build_whitelist(extras={"cholecalciferol"})
    text = "The patient was given cholecalciferol daily."
    findings = spell_check.spell_check_paragraph(text, sc, whitelist=whitelist)
    assert not any(f.original.lower() == "cholecalciferol" for f in findings), (
        f"medical term was flagged: {[f.original for f in findings]}"
    )


def test_spell_check_whitelists_morphological_variant_of_dic_term():
    """Plural / -s form of a dic term should also pass via morph variants."""
    sc = spell_check.get_spell_checker(custom_words={"cholecalciferol"})
    if sc is None:
        pytest.skip("pyspellchecker not installed")
    whitelist = spell_check.build_whitelist(extras={"cholecalciferol"})
    text = "Various cholecalciferols were measured."
    findings = spell_check.spell_check_paragraph(text, sc, whitelist=whitelist)
    assert not any(f.original.lower().startswith("cholecalciferol") for f in findings)


def test_spell_check_still_detects_genuine_typo_when_dict_loaded():
    """Loading the medical dict must not disable genuine typo detection."""
    sc = spell_check.get_spell_checker(custom_words={"cholecalciferol", "ibuprofen"})
    if sc is None:
        pytest.skip("pyspellchecker not installed")
    whitelist = spell_check.build_whitelist(extras={"cholecalciferol", "ibuprofen"})
    text = "Please recieve the shipment and administer ibuprofen."
    findings = spell_check.spell_check_paragraph(text, sc, whitelist=whitelist)
    typos = [f.original for f in findings]
    assert "recieve" in typos
    # And the whitelisted drug name is untouched.
    assert "ibuprofen" not in typos


def test_spell_check_bundled_medspl800_integration():
    """End-to-end: the bundled MedSpl800 dic shields a sampled term, and a
    genuine misspelling in the same paragraph is still flagged."""
    custom_dict.load_all_custom_dicts.cache_clear()
    try:
        dic_words = custom_dict.load_all_custom_dicts()
    finally:
        custom_dict.load_all_custom_dicts.cache_clear()
    if not dic_words:
        pytest.skip("no bundled custom dicts available in this environment")
    # Sanity: the bundle actually loaded many terms, not just a stub.
    assert len(dic_words) > 1000
    # Pick a term we know is in the shipped MedSpl800 list.
    assert "cholecalciferol" in dic_words, (
        "medspl800.dic should contribute 'cholecalciferol'"
    )

    sc = spell_check.get_spell_checker(custom_words=dic_words)
    if sc is None:
        pytest.skip("pyspellchecker not installed")
    whitelist = spell_check.build_whitelist(extras=dic_words)
    text = "The dose of cholecalciferol was adjusted but we recieve more data."
    findings = spell_check.spell_check_paragraph(text, sc, whitelist=whitelist)
    flagged = [f.original for f in findings]
    assert "cholecalciferol" not in flagged
    assert "recieve" in flagged


# ---------------------------------------------------------------------------
# GP009 hyphen-stripping rule — documents the shape of the suggestion so the
# frontend's `applyCasePreservingStructure` has a stable target.
# ---------------------------------------------------------------------------

def _run_profile_on(text: str, profile_key: str):
    """Run the full engine against `text` using one of the bundled profiles."""
    import json
    from pathlib import Path
    from app.processing.language_editing import engine, segmenter, rules

    profile_path = (
        Path(__file__).resolve().parent.parent
        / "app"
        / "processing"
        / "language_editing"
        / "config"
        / "profiles"
        / f"{profile_key}.json"
    )
    with profile_path.open("r", encoding="utf-8") as fh:
        profile = json.load(fh)
    rs = rules.load_rules_from_dict(profile.get("rules", []))
    hs = rules.load_house_style_from_dict(profile.get("variant_to_canonical", {}))
    sents = segmenter.sentences(text)
    return engine.analyze(text, rs, hs, sents)


@pytest.mark.parametrize("profile_key", ["uk", "us"])
@pytest.mark.parametrize(
    "word,expected_suggestion",
    [
        ("non-nurturing", "nonnurturing"),
        ("anti-anxiety", "antianxiety"),
        ("pre-intervention", "preintervention"),
        ("post-intervention", "postintervention"),
        ("multi-step", "multistep"),
        ("sub-field", "subfield"),
    ],
)
def test_gp009_strips_hyphen_in_suggestion(profile_key, word, expected_suggestion):
    """GP009 (`non|pre|post|anti|multi|sub` + hyphen + base) emits a
    hyphen-free suggestion. The frontend's popover case controls are what
    restore the hyphen when the user applies UPPERCASE / lowercase / Title
    Case — this test locks in the backend shape they consume."""
    text = f"The {word} environment was studied."
    findings = _run_profile_on(text, profile_key)
    gp009 = [f for f in findings if f.rule_id == "GP009" and f.original == word]
    assert gp009, (
        f"GP009 did not fire on {word!r} with profile {profile_key!r}; "
        f"findings: {[(f.rule_id, f.original) for f in findings]}"
    )
    f = gp009[0]
    assert f.suggestion == expected_suggestion
    # Original keeps the hyphen; the frontend uses this field to re-hydrate
    # structure when the user asks for a case transform.
    assert "-" in f.original
    # And the finding is categorised as grammar / house-style, not as a typo.
    assert f.category == "grammar"


def test_spell_check_preserves_case_of_document():
    """The dic is lowercased internally but documents keep their own case."""
    sc = spell_check.get_spell_checker(custom_words={"cholecalciferol"})
    if sc is None:
        pytest.skip("pyspellchecker not installed")
    whitelist = spell_check.build_whitelist(extras={"cholecalciferol"})
    # Mid-sentence lowercase occurrence is the one that would otherwise fire;
    # title-case / all-caps are already skipped by the checker as proper nouns
    # or acronyms. We only verify no finding is emitted (so no replacement is
    # suggested that would mutate case).
    text = "Daily cholecalciferol levels were normal."  # Title-case start irrelevant
    findings = spell_check.spell_check_paragraph(text, sc, whitelist=whitelist)
    assert not any(f.original.lower() == "cholecalciferol" for f in findings)


