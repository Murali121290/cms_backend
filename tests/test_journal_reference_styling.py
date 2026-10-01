"""Local reference structuring: bib_* on the reference list and cite_bib on in-text citations."""
import docx
from docx.enum.style import WD_STYLE_TYPE

from app.domains.journals.production import apply_local_reference_styles
from app.processing.local_reference_styler import reference_spans

REFS = [
    "1. Epstein RM, Hundert EM. Defining and assessing professional competence. JAMA. Jan 9, "
    "2002;287(2):226-235. doi: 10.1001/jama.287.2.226 Medline: 11779266",
    "2. Hoy D, March L, Brooks P, et al. The global burden of low back pain: estimates from the Global "
    "Burden of Disease 2010 study. Ann Rheum Dis. Jun 2014;73(6):968-974. doi: 10.1136/annrheumdis-2013-204428",
]


def test_jmir_bracketed_doi_and_medline():
    text = ("1.\tMergen M, Graf N, Meyerheim M. Reviewing virtual reality in medical education. BMC Med Educ. "
            "Jul 23, 2024;24(1):788. [doi: 10.1186/s12909-024-05777-5] [Medline: 39044186]")
    fields = {st: text[s:e] for s, e, st in reference_spans(text)}
    assert fields["bib_doi"] == "10.1186/s12909-024-05777-5"
    assert fields["bib_medline"] == "39044186"
    assert fields["bib_journal"] == "BMC Med Educ" and fields["bib_fpage"] == "788"


def test_reference_with_hyperlink_is_styled_and_link_text_kept(tmp_path):
    """A linked "Medline: n" (w:hyperlink) must neither skip the reference nor vanish from the XHTML."""
    import copy
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.processing.docx_to_xhtml_runs import DocxToXhtmlRunsEngine

    path = str(tmp_path / "ms.docx")
    doc = docx.Document()
    doc.styles.add_style("Reference-Numbered", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("References", style="Heading 1")
    p = doc.add_paragraph("1. Epstein RM, Hundert EM. Defining competence. JAMA. Jan 9, 2002;287(2):226-235. [",
                          style="Reference-Numbered")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), doc.part.relate_to("https://pubmed.ncbi.nlm.nih.gov/11779266/",
                                             "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                             is_external=True))
    run = copy.deepcopy(p.runs[0]._r)
    run.find(qn("w:t")).text = "Medline: 11779266"
    link.append(run)
    p._p.append(link)
    p.add_run("]")
    doc.save(path)

    # a HYPERLINK field around the DOI too: begin / instrText / separate / result / end
    def _field_run(kind=None, instr=None, text=None):
        r = OxmlElement("w:r")
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r.append(fc)
        if instr:
            it = OxmlElement("w:instrText"); it.text = instr; r.append(it)
        if text:
            t = OxmlElement("w:t"); t.text = text; r.append(t)
        return r
    p.add_run(" [")
    for r in (_field_run("begin"), _field_run(instr=' HYPERLINK "https://doi.org/10.1001/jama.287.2.226" '),
              _field_run("separate"), _field_run(text="doi: 10.1001/jama.287.2.226"), _field_run("end")):
        p._p.append(r)
    p.add_run("]")
    doc.save(path)
    before = docx.Document(path).paragraphs[1].text

    result = apply_local_reference_styles(path, "brackets")
    assert result["references_styled"] == 1 and result["hyperlinks_removed"] == 2
    para = docx.Document(path).paragraphs[1]
    assert para.text == before  # link text kept
    assert not para._p.findall(qn("w:hyperlink")) and not para._p.findall(".//" + qn("w:instrText"))
    xhtml = DocxToXhtmlRunsEngine().convert(path)
    assert "Medline: " in xhtml and "11779266" in xhtml
    assert "<li>" not in xhtml  # a reference style is not rendered as an HTML list ("1. 1. ...")


def _blocks(texts_styles):
    from app.domains.journals.manuscript import Block, categorize_style
    return [Block(idx=i, style=s, category=categorize_style(s), text=t) for i, (t, s) in enumerate(texts_styles)]


def test_only_ref_n_and_ref_u_between_markers_are_references():
    from app.domains.journals.checks.references import find_reference_blocks
    blocks = _blocks([("Intro text [1].", "TXT"), ("References", "H1"), ("<ref-open>", "REF-OPEN"),
                      ("1. Smith J. A. J. 2020;1:1-2.", "REF-N"), ("Note inside", "TXT"), ("Brown K. (2019). B. C, 2, 3.", "REF-U"),
                      ("<ref-close>", "REF-CLOSE"), ("MR: mixed reality", "REF-U")])
    assert [b.text[:5] for b in find_reference_blocks(blocks)] == ["1. Sm", "Brown"]


def test_without_markers_the_zone_ends_at_back_matter():
    from app.domains.journals.checks.references import find_reference_blocks
    blocks = _blocks([("References", "H1"), ("1. Smith J. A. J. 2020;1:1-2.", "REF-N"), ("2. Lee K. B. J. 2021;2:3-4.", "REF-N"),
                      ("Abbreviations", "REF-U"), ("MR: mixed reality", "REF-U"), ("Please cite as:", "REF-U")])
    assert [b.text[:2] for b in find_reference_blocks(blocks)] == ["1.", "2."]


def _manuscript(path):
    doc = docx.Document()
    for name in ("REF-OPEN", "REF-N-MID"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("Assessment is critical to certify clinical competence [1]. It recurs (2019) [1,2].")
    doc.add_paragraph("References", style="Heading 1")
    doc.add_paragraph("<ref-open>", style="REF-OPEN")
    for text in REFS:
        doc.add_paragraph(text, style="REF-N-MID")
    doc.save(path)


def _styled(para, prefix):
    return [(r.text, r.style.name) for r in para.runs if r.style.name.startswith(prefix)]


def test_jmir_date_form_gets_journal_volume_and_pages():
    fields = {st: REFS[0][s:e] for s, e, st in reference_spans(REFS[0])}
    assert fields["bib_journal"] == "JAMA"
    assert fields["bib_year"] == "2002"
    assert (fields["bib_volume"], fields["bib_issue"], fields["bib_fpage"], fields["bib_lpage"]) == ("287", "2", "226", "235")
    assert fields["bib_doi"] == "10.1001/jama.287.2.226"


def test_references_and_citations_are_styled_without_changing_text(tmp_path):
    path = str(tmp_path / "ms.docx")
    _manuscript(path)
    before = [p.text for p in docx.Document(path).paragraphs]

    result = apply_local_reference_styles(path, "brackets")

    doc = docx.Document(path)
    assert [p.text for p in doc.paragraphs] == before
    assert result["references_styled"] == 2  # the <ref-open> marker is not a reference
    assert result["citations_styled"] == 2
    body = doc.paragraphs[1]
    assert _styled(body, "cite_") == [("[1]", "cite_bib"), ("[1,2]", "cite_bib")]  # "(2019)" untouched
    ref1 = dict((st, t) for t, st in _styled(doc.paragraphs[4], "bib_"))
    assert ref1["bib_journal"] == "JAMA" and ref1["bib_article"] == "Defining and assessing professional competence"

    again = apply_local_reference_styles(path, "brackets")  # idempotent
    assert again["references_styled"] == 0 and again["citations_styled"] == 0
    assert [p.text for p in docx.Document(path).paragraphs] == before
